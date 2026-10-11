"""Layout/guest replay regressions: fictional candidates, fake models, local PDF/SQL."""
from copy import deepcopy
import io
import json
from pathlib import Path
import re
import sys
import unittest
from unittest.mock import AsyncMock, patch

from lxml.html import fromstring
from docx import Document
from langchain_core.messages import AIMessage
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'backend'))
from app.agents import resume_graph as graph
from app.agents.resume_writer import run_resume_writer, _STRUCTURE_CONTRACT, _WRITE_CONTRACT, _ONE_PASS_CONTRACT
from app.agents.triage_chain import TriageAction, TriageDecision
from app.config import get_settings
from app.database import Base
from app.schemas.editor import ResumeData
from app.services.data_crypto import DataCrypto
from app.services.model_privacy import CONTACT_PRIVACY_NOTICE, ContactRedactor, redact_model_messages
from app.services.profile_bridge import sanitize_fact_history
from app.services.resume_field_hygiene import strip_pr_prefix
from app.services.renderer import list_templates, render_resume
from app.services.resume_pdf import count_resume_pdf_pages
from app.services.resume_review import resume_to_text
from app.services.session_store import load_resume_session
from app.services.structure_plan import apply_plan_structure, plan_guidance, plan_summary, select_plan
from app.tools.resume_data_docx import resume_data_to_docx
from test_replay_plan_intake_regressions import PlanModel
from test_structure_plan import card, source, writer_draft


PROFILE = {'role_family': 'pm', 'requirements': [
    {'id': 'r1', 'text': '评测方案与 badcase 分析'}, {'id': 'r2', 'text': '用户需求分析'}],
    'core_signals': [{'id': 's2', 'signal': '评测方案与 badcase 分析'},
                     {'id': 's3', 'signal': '需求拆解'}, {'id': 's5', 'signal': '反馈分类'}]}
PERSONAL = ('姓名林小舟，人在深圳，求职意向是 AI 产品实习。\n'
            '电话 13800000000，邮箱 fiction@candidate.test。\n'
            '教育：哈尔滨商业大学，计算机科学与技术本科，2023.09–2027.06，获得过校级三等奖学金。')


def select(root, selector):
    if selector.startswith('.'):
        matches = root.find_class(selector[1:])
    else:
        path = selector.split('"')[1]
        matches = root.xpath('.//*[@data-field-path=$path]', path=path)
    return matches[0] if matches else None


def layout_fixture():
    return ResumeData.model_validate({
        'basic_info': {'name': '合成候选人', 'objective': 'AI 产品实习', 'email': 'fiction@candidate.test'},
        'summary': '通过用户反馈与开源修复展示问题分析能力',
        'education': [{'school': '示例大学', 'major': '计算机科学与技术', 'degree': '本科',
                       'start_date': '2023.09', 'end_date': '2027.06', 'highlights': '校级三等奖学金'}],
        'projects': [{'name': f'虚构项目{i}', 'polished_bullets': [
            '整理用户反馈并分类，形成问题清单和迭代建议供小组讨论。',
            '围绕核心流程制作原型，记录测试样例并复查反馈结果。',
            '对比不同方案的输出，记录失败样例和已有功能的限制。']} for i in range(1, 4)],
        'open_source': [
            {'name': 'VoiceMem', 'pr': 'PR #10', 'bullets': ['PR #10：修复跨语言记忆污染']},
            {'name': 'OpenSquilla', 'summary': '7k★ 的 AI Agent 项目', 'pr': 'PR #21',
             'bullets': ['PR #21：修复空输入的参数校验']}],
        'awards': [{'name': f'虚构竞赛{i}', 'level': '省级二等奖',
                    'bullets': [f'虚构竞赛{i}。']} for i in range(1, 6)],
        'skill_groups': {'开发工具': ['Python', 'Java', 'SQL'], '产品方法': ['需求分析', '原型设计']},
    }).model_dump()


class LayoutRegression(unittest.TestCase):
    def test_all_six_templates_keep_awards_in_one_paragraph(self):
        data = layout_fixture()
        data['awards'][1]['bullets'].append('负责现场演示和答辩')
        data['awards'][0]['hidden'] = True
        before = deepcopy(data)
        for template in list_templates():
            with self.subTest(template=template['id']):
                soup = fromstring(render_resume(template['id'], data))
                paragraph = select(soup, '.awards-compact')
                self.assertIsNotNone(paragraph)
                text = paragraph.text_content().strip()
                self.assertEqual(text.count('；'), 3)
                self.assertNotIn('虚构竞赛1', text)
                self.assertIn('虚构竞赛2 省级二等奖', text)
                self.assertIn('负责现场演示和答辩', soup.text_content())
                self.assertIsNone(select(soup, '[data-field-path="/awards/1/bullets/0"]'))
                self.assertIsNotNone(select(soup, '[data-field-path="/awards/1/bullets/1"]'))
        self.assertEqual(data, before)

    def test_word_awards_are_one_paragraph_with_only_substantive_explanation(self):
        data = layout_fixture()
        data['awards'][1]['bullets'].append('负责现场演示和答辩')
        paragraphs = [p.text for p in Document(io.BytesIO(resume_data_to_docx(data))).paragraphs]
        award_lines = [text for text in paragraphs if '虚构竞赛' in text]
        self.assertEqual(sum('；' in text for text in award_lines), 1)
        self.assertEqual(award_lines[0].count('；'), 4)
        self.assertEqual(sum('负责现场演示和答辩' in text for text in paragraphs), 1)

    def test_positioning_is_header_only_in_six_templates_and_word(self):
        data = layout_fixture()
        data['self_evaluation'] = '愿意认真记录问题并复查自己的工作'
        data['projects'] = data['projects'][:2]
        for template in list_templates():
            with self.subTest(template=template['id']):
                soup = fromstring(render_resume(template['id'], data))
                header = select(soup, '.header-text')
                self.assertEqual(select(header, '[data-field-path="/summary"]').text_content(), data['summary'])
                self.assertEqual(soup.text_content().count(data['summary']), 1)
                self.assertEqual(select(soup, '[data-field-path="/self_evaluation"]').text_content(), data['self_evaluation'])
                self.assertEqual(count_resume_pdf_pages(template['id'], data), 1)
        paragraphs = [p.text for p in Document(io.BytesIO(resume_data_to_docx(data))).paragraphs]
        self.assertLess(paragraphs.index(data['summary']), paragraphs.index('教育背景'))
        self.assertEqual(paragraphs.count(data['summary']), 1)
        self.assertEqual(paragraphs[paragraphs.index('自我评价') + 1], data['self_evaluation'])

    def test_no_positioning_in_self_evaluation_and_schema_preserves_header(self):
        data = apply_plan_structure(layout_fixture(), select_plan(card()))
        self.assertEqual(data['summary'], card()['recommended']['positioning'])
        self.assertFalse(data['self_evaluation'])
        self.assertEqual(ResumeData.model_validate(data).model_dump()['summary'], data['summary'])
        for template in list_templates():
            self.assertNotIn('>自我评价<', render_resume(template['id'], data))

    def test_repository_summary_and_pr_share_title_and_bullets_drop_only_matching_prefix(self):
        data = layout_fixture()
        data['open_source'][0]['bullets'] += ['PR #11：另一次修复', '在 PR #10 中补充测试']
        for template in list_templates():
            with self.subTest(template=template['id']):
                soup = fromstring(render_resume(template['id'], data))
                title = select(soup, '[data-field-path="/open_source/1/name"]').getparent().text_content().strip()
                self.assertIn('OpenSquilla（7k★ 的 AI Agent 项目）', title)
                self.assertIn('PR #21', title)
                self.assertEqual(select(soup, '[data-field-path="/open_source/0/bullets/0"]').text_content(), '修复跨语言记忆污染')
                self.assertIn('PR #11：另一次修复', soup.text_content())
                self.assertIn('在 PR #10 中补充测试', soup.text_content())
        paragraphs = [p.text for p in Document(io.BytesIO(resume_data_to_docx(data))).paragraphs]
        self.assertIn('OpenSquilla（7k★ 的 AI Agent 项目）｜ PR #21', paragraphs)
        self.assertIn('修复跨语言记忆污染', paragraphs)
        self.assertNotIn('PR #10：修复跨语言记忆污染', resume_to_text(data))

    def test_t001_three_projects_two_prs_five_awards_fit_one_page(self):
        self.assertEqual(count_resume_pdf_pages('t001-jianyue', layout_fixture()), 1)

    def test_guidance_and_bounded_summary_replace_requirement_and_signal_ids(self):
        planned = card()
        planned['recommended']['thesis'] = '围绕 r1、r2 展示实践，对应 s2/s3/s5 和 r99'
        planned['recommended']['experience_plan'][0]['reason'] = '修复跨语言记忆污染，命中 r1、r2，对应 s2/s3/s5。' * 5
        planned['alternatives'][0]['thesis'] = '以 s2 展示 r2'
        guidance = plan_guidance(planned, source(), PROFILE)
        prose = json.dumps(guidance['structure_plan'], ensure_ascii=False)
        # Machine identity slots stay intact; every prose value must be readable.
        for row in [guidance['structure_plan']['recommended'], *guidance['structure_plan']['alternatives']]:
            for text in [row['thesis'], row['positioning'], *[entry['reason'] for entry in row['experience_plan']]]:
                self.assertIsNone(re.search(r'\b[rs]\d+\b', text))
        self.assertIn('评测方案与 badcase 分析', prose)
        summary = plan_summary(select_plan(planned), source(), PROFILE)
        self.assertLessEqual(len('**这版的思路**：' + summary), 150)
        self.assertIsNone(re.search(r'\b[rs]\d+\b', summary))
        self.assertIn('EchoKit', summary)

    def test_privacy_notice_explains_saved_contacts_and_public_acknowledgement(self):
        self.assertIn('系统保存', CONTACT_PRIVACY_NOTICE)
        self.assertIn('电话、邮箱已记下', CONTACT_PRIVACY_NOTICE)
        messages = redact_model_messages([AIMessage(content=PERSONAL)], ContactRedactor())
        outbound = str(messages)
        self.assertNotIn('13800000000', outbound)
        self.assertNotIn('fiction@candidate.test', outbound)
        self.assertIn('不提占位符', outbound)

    def test_contracts_require_human_readable_reasons_and_pr_in_header(self):
        self.assertIn('要求原文', _STRUCTURE_CONTRACT)
        for contract in (_WRITE_CONTRACT, _ONE_PASS_CONTRACT):
            self.assertIn('PR 号在标题行', contract)
            self.assertIn('bullet 不重复', contract)

    def test_pr_prefix_variants_preserve_other_prs_and_in_sentence_references(self):
        for prefix in ('PR #10：', 'pr#10: ', 'PR #10 ｜ ', '#10：', 'PR #10 '):
            with self.subTest(prefix=prefix):
                self.assertEqual(strip_pr_prefix(prefix + '修复跨语言记忆污染', 'PR #10'), '修复跨语言记忆污染')
        for text in ('PR #100：另一项修复', 'PR #11：另一项修复', '在 PR #10 中补充测试', '10 条回归测试'):
            self.assertEqual(strip_pr_prefix(text, 'PR #10'), text)

    def test_three_lead_names_and_twenty_character_reasons_fit_150_characters(self):
        planned = select_plan(card())
        planned['thesis'] = '围绕已有经历说明分析问题和完成验证的过程。' * 12
        planned['experience_plan'] = [{'item_id': iid, 'treatment': 'lead', 'reason': '解释为什么重点写这段经历以及如何支撑岗位要求' * 6}
                                      for iid in ('p1', 'o1', 'a1')]
        summary = plan_summary(planned, source(), PROFILE)
        self.assertLessEqual(len('**这版的思路**：' + summary), 150)
        for name in ('纸舟计划', 'EchoKit', '合成竞赛'):
            self.assertIn(name, summary)
        self.assertEqual(len(re.findall('（([^）]*)）', summary)), 3)
        self.assertTrue(all(len(reason) <= 20 for reason in re.findall('（([^）]*)）', summary)))

    def test_corrected_major_survives_while_superseded_major_and_old_versions_are_cleared(self):
        old = {'id': 'old-major', 'field': 'major', 'status': 'superseded',
               'text': '计算机科学与技术本科', 'source_quote': '计算机科学与技术本科'}
        active = {'id': 'new-major', 'field': 'major', 'status': 'active',
                  'text': '计算机科学与技术', 'source_quote': '计算机科学与技术'}
        state = {'collected_info': {'education': [{'item_id': 'edu', 'facts': [old, active]}]},
                 'resume_data': {'education': [{'item_id': 'edu', 'major': active['text']}]},
                 'resume_versions': [{'resume_data': {'education': [{'item_id': 'edu', 'major': old['text']}]}}]}
        result = sanitize_fact_history(state)
        self.assertEqual(result['resume_data']['education'][0]['major'], active['text'])
        self.assertEqual(result['resume_versions'][0]['resume_data']['education'][0]['major'], '')
        state['collected_info']['education'][0]['facts'][1]['status'] = 'deleted'
        self.assertEqual(sanitize_fact_history(state)['resume_data']['education'][0]['major'], '')

    def test_header_summary_has_own_version_change_and_stale_fact_invalidation(self):
        from app.services.resume_versions import document_diff, text_hash
        data = {'summary': '修复跨语言记忆污染', 'self_evaluation': '做事认真'}
        self.assertEqual(document_diff({}, data)[0]['section'], 'summary')
        collected = source()
        collected['open_source'][0]['facts'][0]['status'] = 'deleted'
        cleaned = sanitize_fact_history({'collected_info': collected, 'resume_data': data,
            'resume_meta': {'summary': {'fact_ids': ['f2'], 'text_hash': text_hash(data['summary'])}}})
        self.assertEqual(cleaned['resume_data']['summary'], '')
        self.assertEqual(cleaned['resume_data']['self_evaluation'], '做事认真')
        self.assertNotIn('summary', cleaned['resume_meta'])


class WriterLayoutRegression(unittest.IsolatedAsyncioTestCase):
    async def test_expanded_positioning_is_guarded_before_becoming_header_summary(self):
        selected = select_plan(card())
        selected['positioning'] = '对应 r1 的实践者'
        model = PlanModel([('WriterDraft', writer_draft())], fail_ids=['plan_positioning_0'])
        result = await run_resume_writer(source(), PROFILE, model, phase='write', previous_plan=selected)
        checked = next(row for row in model.batches[0] if row['id'] == 'plan_positioning_0')
        self.assertEqual(checked['text'], '对应 评测方案与 badcase 分析 的实践者')
        self.assertEqual(result['resume_data']['summary'], '')
        self.assertFalse(result['confirmed_plan']['positioning'])

    async def test_model_plan_prose_uses_requirement_text_but_machine_references_stay_intact(self):
        planned = card()
        planned['recommended']['thesis'] += '，对应 r1 和 s2'
        planned['recommended']['experience_plan'][0]['reason'] += '命中 r1、r2，对应 s2/s3/s5；r99。'
        result = await run_resume_writer(source(), PROFILE, PlanModel([('StructurePlan', planned)]), phase='plan')
        selected = result['structure_plan']['recommended']
        self.assertIn('评测方案与 badcase 分析', selected['experience_plan'][0]['reason'])
        self.assertIsNone(re.search(r'(?<![A-Za-z0-9_])[rs]\d+(?![A-Za-z0-9_])', selected['thesis']))
        self.assertEqual(selected['requirements'][0]['req_id'], 'r1')
        self.assertEqual(selected['experience_plan'][0]['item_id'], 'o1')

    async def test_postprocessing_removes_duplicate_pr_and_rehashes_bullet_metadata(self):
        collected = source()
        collected['open_source'][0]['facts'].extend([
            {'id': 'pr', 'text': 'PR #10', 'source_quote': 'PR #10', 'field': 'pr', 'status': 'active'}])
        draft = writer_draft()
        draft['items'][0]['bullets'][0]['text'] = 'PR #10：修复跨语言记忆污染'
        draft['items'][0]['bullets'][0]['fact_ids'].append('pr')
        for phase in ('write', 'full'):
            with self.subTest(phase=phase):
                selected = select_plan(card())
                match = {key: selected[key] for key in ('requirements', 'experience_plan', 'section_order')}
                response = ('WriterDraft', draft) if phase == 'write' else ('OnePassOutput', {'match_plan': match, 'writer': draft})
                model = PlanModel([response])
                written = await run_resume_writer(collected, PROFILE, model, phase=phase,
                                                  previous_plan=selected if phase == 'write' else None)
                row = written['resume_data']['open_source'][0]
                self.assertEqual(row['bullets'], ['修复跨语言记忆污染'])
                from app.services.resume_versions import text_hash
                self.assertEqual(written['resume_meta']['bullets']['o1'][0]['text_hash'], text_hash(row['bullets'][0]))

    async def test_confirmed_write_preserves_user_self_evaluation_separately(self):
        collected = source()
        collected['personal_summary'] = '愿意认真记录问题并复查自己的工作'
        written = await run_resume_writer(collected, PROFILE, PlanModel([('WriterDraft', writer_draft())]),
                                          phase='write', previous_plan=select_plan(card()))
        self.assertEqual(written['resume_data']['summary'], card()['recommended']['positioning'])
        self.assertEqual(written['resume_data']['self_evaluation'], collected['personal_summary'])

    async def test_full_guest_third_turn_intake_plan_direct_write_and_sql_save(self):
        engine = create_engine('sqlite://', connect_args={'check_same_thread': False}, poolclass=StaticPool)
        Base.metadata.create_all(engine)
        factory = sessionmaker(bind=engine)
        self.addCleanup(engine.dispose)
        model = PlanModel([('StructurePlan', card()), ('WriterDraft', writer_draft())])
        compiled = graph.build_graph(model)
        session = 'fictional-layout-guest'
        initial = {**graph._init_state(), 'stage': 'COLLECTING', 'collected_info': source(),
                   'candidate_history_reconciled': True, 'jd_info': {'job_title': 'AI 产品实习', 'role_profile': PROFILE}}
        from langchain_core.messages import HumanMessage
        initial['messages'] = [HumanMessage(content='我想找 AI 产品实习'), AIMessage(content='说说经历'),
                               HumanMessage(content='我在纸舟计划整理用户反馈并分类'), AIMessage(content='请补充教育和联系方式')]
        await compiled.aupdate_state({'configurable': {'thread_id': session}}, initial)
        from test_replay_bug_regressions import ExtractModel
        extractor = ExtractModel({'items': [{'section': 'education', 'name': '哈尔滨商业大学', 'facts': [
            {'field': 'degree', 'text': '本科', 'source_quote': '本科', 'kind': 'context'},
            {'field': 'major', 'text': '计算机科学与技术本科', 'source_quote': '计算机科学与技术本科', 'kind': 'context'}]}]})
        with patch.object(get_settings(), 'LLM_PLAN_CONFIRM', True), \
             patch.object(graph, 'get_graph', return_value=compiled), \
             patch.object(graph, '_make_non_streaming_llm', return_value=extractor), \
             patch.object(graph, 'run_triage', AsyncMock(return_value=TriageDecision(action=TriageAction.ASK_EXPERIENCE, reason='fake'))), \
             patch.object(graph, 'stream_public_reply', AsyncMock(return_value=AIMessage(content='电话、邮箱已记下'))), \
             patch('app.services.session_store.SessionLocal', factory), \
             patch('app.services.data_crypto.get_crypto', return_value=DataCrypto({'test': b'x' * 32}, 'test')), \
             patch('app.services.profile_bridge.archive_event_for_message', return_value=None):
            async def turn(text):
                return ''.join([chunk async for chunk in graph.stream_agent([{'role': 'user', 'content': text}], session, model)])
            await turn(PERSONAL)
            saved = load_resume_session(session)
            intake_major = saved['resume_data']['education'][0]['major']
            await turn('请生成简历')
            snapshot = (await compiled.aget_state({'configurable': {'thread_id': session}})).values
            self.assertTrue(snapshot['pending_plan'])
            reply = await turn('先出一版看看')
            saved = load_resume_session(session)
        self.assertEqual(len(extractor.calls), 1)
        self.assertTrue(saved['resume_meta']['confirmed_plan'])
        self.assertTrue(saved['resume_versions'])
        edu = saved['resume_data']['education'][0]
        self.assertEqual((edu['school'], edu['major'], edu['degree']), ('哈尔滨商业大学', '计算机科学与技术', '本科'))
        self.assertEqual(intake_major, '计算机科学与技术')
        self.assertEqual((edu['start_date'], edu['end_date']), ('2023.09', '2027.06'))
        self.assertIn('校级三等奖学金', edu['highlights'])
        self.assertEqual(count_resume_pdf_pages('t001-jianyue', saved['resume_data']), 1)
        self.assertEqual(reply.count('**这版的思路**'), 1)
        self.assertNotIn('方案摘要', reply)
        thought = reply.split('**这版的思路**：', 1)[1].split('\n\n', 1)[0]
        self.assertLessEqual(len('**这版的思路**：' + thought), 150)


if __name__ == '__main__':
    unittest.main()
