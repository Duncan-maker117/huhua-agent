"""Web-stored documents round-trip through the official MCP HTTP client, offline."""

import base64
from copy import deepcopy
import io
import json
import unittest
from unittest.mock import patch

import pdfplumber
from docx import Document
from pydantic import ValidationError

import test_mcp_protocol as legacy
from app.mcp.server import MCPResume
from app.models.models import ResumeSession
from app.services.conversation_store import create_conversation
from app.services.session_store import save_resume_session


TEMPLATE = 't001-jianyue'
RESUME_TOOLS = ('huhua_resume_measure', 'huhua_resume_preview', 'huhua_resume_save', 'huhua_resume_export')


def web_document():
    return MCPResume.model_validate({
        'basic_info': {'name': '合成用户', 'phone': '13800001234', 'availability': '随时到岗',
                       'links': [{'text': '主页', 'url': 'https://synthetic.example.test'}]},
        'education': [{'school': '合成大学', 'major': '计算机', 'highlights': '完成课程设计'}],
        'work_experience': [{'company': '合成公司', 'title': '开发', 'employment_type': 'part_time',
                             'bullets': ['维护内部工具']}],
        'projects': [{'name': '合成项目', 'positioning': '课程目录', 'project_type': '课程项目',
                      'polished_bullets': ['制作课程目录原型'],
                      'links': [{'text': '演示', 'url': 'https://demo.example.test'}]},
                     {'name': '保留的隐藏项目', 'hidden': True, 'description': '不显示但须保留'}],
        'campus': [{'name': '合成社团', 'bullets': ['整理活动资料']}],
        'open_source': [{'name': '合成开源项目', 'summary': '修订说明文档'}],
        'awards': [{'name': '合成奖项', 'level': '校级'}],
        'skills': ['熟悉 Python'], 'skill_groups': {'技术/工具~基础': ['熟悉 Python']},
        'certificates': ['合成证书'], 'summary': '寻找开发岗位', 'self_evaluation': '愿意核对细节',
        'experience_section_title': '兼职经历', 'verification_notes': ['网页端待核对事项'],
        'layout': {'density': 1, 'density_source': 'user'},
        'sections': [{'id': 'projects', 'kind': 'projects', 'title': '我的项目'},
                     {'id': 'skills', 'kind': 'skills', 'title': '专业技能'},
                     {'id': 'custom', 'kind': 'custom', 'title': '补充说明', 'content': '保留自定义正文'}],
    }).model_dump()


class MCPRoundTripTests(unittest.IsolatedAsyncioTestCase):
    asyncSetUp = legacy.MCPProtocolTests.asyncSetUp
    cleanup = legacy.MCPProtocolTests.cleanup
    client = legacy.MCPProtocolTests.client
    call = legacy.MCPProtocolTests.call
    add = legacy.MCPProtocolTests.add

    def store_web_draft(self):
        document = web_document()
        stored = deepcopy(document)
        stored.update(resume_strategy={'mode': 'capability', 'internal': 'strategy-private'},
                      resume_meta={'source': 'writer'}, future_internal={'trace': 'server-private'})
        stored['basic_info']['internal_note'] = 'basic-private'
        stored['projects'][0]['fact_ids'] = ['server-fact-id']
        stored['projects'][0]['provenance'] = {'source': 'server-private'}
        stored['projects'][0]['links'][0]['internal_source'] = 'link-private'
        stored['sections'][2]['internal_note'] = 'section-private'
        stored['layout']['measured_pages'] = 1
        session = create_conversation('account:a')['session_id']
        # The same persistence boundary used by the web chat, with an isolated DB.
        save_resume_session(session, stored, stage='EDITING', resume_meta={'source': 'writer'})
        with self.factory() as db:
            row = db.get(ResumeSession, session)
            before = deepcopy(row.resume_data), row.version
        self.assertIn('resume_strategy', before[0])
        self.assertIn('provenance', before[0]['projects'][0])
        return session, document, before

    async def test_web_get_measure_preview_save_export_roundtrip(self):
        session, document, before = self.store_web_draft()
        self.add(category='basic', subject='基本信息', field='name', value='合成用户')
        self.add(category='basic', subject='基本信息', field='phone', value='13800001234')
        with patch('app.services.fact_guard.verify_semantics', side_effect=AssertionError('no model')):
            async with self.client() as client:
                result = await self.call(client, 'huhua_resume_get', {'conversation_id': session})
                latest = await self.call(client, 'huhua_resume_get')
                self.assertEqual(result, latest)
                resume = result['untrusted_text']['resume']
                document['basic_info']['phone'] = '（联系方式已隐藏）'
                document['basic_info']['links'] = []
                self.assertEqual(resume, document)
                MCPResume.model_validate(resume)
                untouched = deepcopy(resume)
                args = {'resume': resume, 'template_id': TEMPLATE}
                measured = await self.call(client, 'huhua_resume_measure', args)
                self.assertGreaterEqual(measured['pages'], 1)
                preview = await client.call_tool('huhua_resume_preview', args)
                self.assertFalse(preview.is_error, preview.content)
                png = next(row for row in preview.content if row.type == 'image')
                self.assertTrue(base64.b64decode(png.data).startswith(b'\x89PNG'))
                self.assertTrue(preview.structured_content['contacts_hidden'])
                saved = await self.call(client, 'huhua_resume_save', args)
                self.assertTrue(saved['saved'])
                self.assertNotEqual(saved['conversation_id'], session)
                with self.factory() as db:
                    new = db.get(ResumeSession, saved['conversation_id']).resume_data
                    for field in ('sections', 'skill_groups', 'layout', 'summary', 'certificates',
                                  'self_evaluation', 'experience_section_title'):
                        self.assertEqual(new[field], resume[field])
                    self.assertTrue(new['projects'][1]['hidden'])
                    self.assertEqual(new['basic_info']['phone'], '13800001234')
                for format in ('pdf', 'docx'):
                    exported = await self.call(client, 'huhua_resume_export', {**args, 'format': format})
                    download = await self.web.get(exported['download_url'])
                    self.assertEqual(download.status_code, 200)
                    content = download.content
                    if format == 'pdf':
                        self.assertTrue(content.startswith(b'%PDF-'))
                        with pdfplumber.open(io.BytesIO(content)) as pdf:
                            text = '\n'.join(page.extract_text() for page in pdf.pages)
                    else:
                        self.assertTrue(content.startswith(b'PK'))
                        doc = Document(io.BytesIO(content))
                        text = '\n'.join(p.text for p in doc.paragraphs)
                        text += '\n'.join(cell.text for table in doc.tables for row in table.rows for cell in row.cells)
                    for visible in ('合成项目', 'Python', '保留自定义正文', '13800001234'):
                        self.assertIn(visible, text)
                    self.assertNotIn('strategy-private', text)
                self.assertEqual(resume, untouched)
                again = await self.call(client, 'huhua_resume_get', {'conversation_id': saved['conversation_id']})
                MCPResume.model_validate(again['untrusted_text']['resume'])
        with self.factory() as db:
            row = db.get(ResumeSession, session)
            self.assertEqual((row.resume_data, row.version), before)

    async def test_get_projection_keeps_authorized_contacts_and_saved_source_tools(self):
        session, document, _ = self.store_web_draft()
        pat = self.tokens.create('a', '合成权限', ['profile:contacts', 'resume:read', 'resume:export'])
        async with self.client(pat) as client:
            result = await self.call(client, 'huhua_resume_get', {'conversation_id': session})
            self.assertEqual(result['untrusted_text']['resume'], document)
            args = {'conversation_id': session, 'template_id': TEMPLATE}
            preview = await client.call_tool('huhua_resume_preview', args)
            self.assertFalse(preview.is_error, preview.content)
            exported = await self.call(client, 'huhua_resume_export', args)
            self.assertEqual((await self.web.get(exported['download_url'])).status_code, 200)

    async def test_invalid_resume_has_safe_pointers_for_every_input_tool(self):
        secret = 'synthetic-private-value-must-not-leak'
        invalid = {'resume_strategy': secret, 'projects': [{'provenance': secret}],
                   'basic_info': {'internal/~note': secret}}
        with self.assertRaises(ValidationError):
            MCPResume.model_validate(invalid)
        async with self.client() as client:
            with patch('app.mcp.server.measure', side_effect=AssertionError('no rendering')):
                for name in RESUME_TOOLS:
                    with self.subTest(tool=name):
                        result = await self.call(client, name, {'resume': invalid, 'template_id': TEMPLATE}, error=True)
                        payload = json.loads(result.content[0].text)
                        self.assertEqual(payload, result.structured_content)
                        self.assertEqual(payload['error'], 'invalid_resume')
                        self.assertEqual(set(payload['fields']), {
                            '/resume_strategy: 不允许的字段', '/projects/0/provenance: 不允许的字段',
                            '/basic_info/internal~1~0note: 不允许的字段'})
                        self.assertNotIn(secret, str(result))
                        self.assertNotIn('ValidationError', str(result))
        rows = self.audit.recent({self.pat['id']})
        self.assertEqual(len(rows), 4)
        self.assertTrue(all(row['status_code'] == 422 for row in rows))
        self.assertNotIn(secret, json.dumps(rows))

    async def test_validation_categories_include_types_ranges_and_argument_paths(self):
        async with self.client() as client:
            for resume, fields in [
                ({'projects': [{'polished_bullets': [123]}]}, ['/projects/0/polished_bullets/0: 需要字符串']),
                ({'layout': {'density': 4}}, ['/layout/density: 超出允许范围']),
                ({'sections': [{'id': 'x', 'kind': 'unknown'}]},
                 ['/sections/0/kind: 不允许的取值', '/sections/0/title: 缺少必填字段']),
                ({'skill_groups': {'技术/工具~基础': [123]}}, ['/skill_groups/技术~1工具~0基础/0: 需要字符串']),
            ]:
                result = await self.call(client, 'huhua_resume_measure',
                                         {'resume': resume, 'template_id': TEMPLATE}, error=True)
                self.assertEqual(result.structured_content['error'], 'invalid_resume')
                self.assertEqual(result.structured_content['fields'], fields)
            result = await self.call(client, 'huhua_resume_preview',
                                     {'resume': {}, 'template_id': TEMPLATE, 'page': 0}, error=True)
            self.assertEqual(result.structured_content['error'], 'invalid_arguments')
            self.assertEqual(result.structured_content['fields'], ['/page: 超出允许范围'])

    async def test_permission_denial_still_precedes_readable_validation(self):
        pat = self.tokens.create('a', '无写入权限', [])
        async with self.client(pat) as client:
            result = await self.call(client, 'huhua_resume_save', {'resume': {'resume_strategy': {}}}, error=True)
            self.assertEqual(result.content[0].text, 'mcp_error_403')
            self.assertIsNone(result.structured_content)

    async def test_sdk_output_validation_is_not_reported_as_client_fields(self):
        secret = 'synthetic-internal-output-never-return'
        async with self.client() as client:
            with patch('app.mcp.server.measure', return_value=[secret]):
                result = await self.call(client, 'huhua_resume_measure',
                                         {'resume': {}, 'template_id': TEMPLATE}, error=True)
            self.assertEqual(result.content[0].text, 'mcp_error_503')
            self.assertIsNone(result.structured_content)
            self.assertNotIn(secret, str(result))

