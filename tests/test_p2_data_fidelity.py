"""B1–B4 regressions use fictional material and local model doubles only."""

import io
import json
from pathlib import Path
import sys
import tempfile
from types import SimpleNamespace
import unittest
from unittest.mock import AsyncMock, patch
import uuid

from docx import Document
from langchain_core.messages import HumanMessage

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "backend"))

from app.agents.fact_extractor import apply_fact_patch, extract_facts_patch
from app.agents.profile_archivist import deterministic_proposals, validate_grounding
from app.agents.resume_graph import _merge_imported_candidate, _serialize_to_resume_data, evaluating_node
from app.agents.resume_writer import run_resume_writer
from app.agents.router import route_intents
from app.agents.turn_actions import absorb_education, choose_edit_plan, overlay_education
from app.schemas.editor import ResumeData
from app.services.candidate_material import parse_existing_resume
from app.services.edit_plan import execute_edit_plan
from app.services.fact_guard import guard_bullets, lexical_fact_errors
from app.services.renderer import list_templates, render_resume
from app.services.resume_facts import active_facts_for_writer, append_editor_facts, new_fact
from app.services.resume_pdf import render_resume_pdf
from app.services.resume_review import resume_to_text
from app.services.resume_sections import get_resume_sections
from app.services.resume_versions import document_diff, valid_bullet_metadata
from app.services.resume_workspace import build_workspace
from app.tools.docx_template_filler import fill_template_bytes
from app.tools.resume_data_docx import resume_data_to_docx


EDUCATION = "示例甲大学，计算机科学与技术本科，2023.09–2027.06，获得过校级三等奖学金"
LISTS = """开源贡献：
- StarBoat 项目（7k★ 的 AI Agent 项目）PR #123：修复缓存键冲突，维护者已合并。
- CloudNote PR #8：修复记忆中的语言混用，已提交等待审核。
竞赛和荣誉：示例创新赛 Top100；示例省大学生程序设计竞赛三等奖；示例校算法赛二等奖
技能：PRD、原型
"""
HEADER = """姓名：测试小青
GitHub 是 https://github.com/fictional-example，个人主页 example.test
可实习时间：每周 6 天，6 个月起
"""


class ExtractModel:
    def __init__(self, value):
        self.value = value
        self.calls = []

    async def ainvoke(self, messages):
        self.calls.append(messages)
        return SimpleNamespace(content=json.dumps(self.value, ensure_ascii=False))


class WriterModel:
    def __init__(self, response):
        self.response = response
        self.calls = []

    def with_structured_output(self, schema):
        return self

    async def ainvoke(self, messages):
        self.calls.append(messages)
        return self.response


def canonical_material():
    facts = parse_existing_resume(HEADER + "教育背景：" + EDUCATION + "\n" + LISTS)
    return facts, _serialize_to_resume_data({"collected_info": facts})


def docx_text(content):
    return "\n".join(paragraph.text for paragraph in Document(io.BytesIO(content)).paragraphs)


class P2ExtractionTests(unittest.IsolatedAsyncioTestCase):
    async def test_p2_all_prs_and_awards_survive_a_first_item_only_model(self):
        quote = "修复缓存键冲突，维护者已合并"
        model = ExtractModel({"items": [{"section": "projects", "name": "StarBoat", "facts": [
            {"kind": "action", "text": quote, "source_quote": quote}]}]})
        sparse = await extract_facts_patch(LISTS, {}, model)
        facts = apply_fact_patch({}, sparse)
        self.assertFalse(facts.get("projects"))
        self.assertEqual([item["name"] for item in facts["open_source"]], ["StarBoat 项目", "CloudNote"])
        self.assertEqual(len(facts["awards"]), 3)
        self.assertIn("等待审核", str(facts["open_source"][1]["facts"]))
        self.assertEqual(len(model.calls), 1)

    async def test_p2_same_repository_multiple_prs_stay_separate_on_reimport(self):
        text = "开源贡献：\n- StarBoat PR #1：修复文件校验。\n- StarBoat PR #2：修复列表分页。"
        imported = parse_existing_resume(text)
        merged = _merge_imported_candidate({}, imported)
        merged = _merge_imported_candidate(merged, imported)
        self.assertEqual(len(merged["open_source"]), 2)
        self.assertEqual({item["pr"] for item in merged["open_source"]}, {"PR #1", "PR #2"})
        sparse = await extract_facts_patch(text, {}, ExtractModel({"items": [{"section": "projects", "name": "StarBoat",
            "facts": [{"kind": "action", "text": "修复文件校验", "source_quote": "修复文件校验"}]}]}))
        self.assertEqual(len(apply_fact_patch({}, sparse)["open_source"]), 2)

    async def test_p2_archive_list_proposals_are_grounded_per_pr(self):
        payload = {"human_text": LISTS, "source_kind": "self"}
        proposals = deterministic_proposals(payload)
        self.assertTrue(proposals)
        for proposal in proposals:
            self.assertTrue(validate_grounding(proposal, payload, {"facts": []}), proposal)

    async def test_p2_extractor_retains_education_field_tags_and_header_links(self):
        message = HEADER + "教育：" + EDUCATION
        model = ExtractModel({"items": [{"section": "education", "name": "示例甲大学", "facts": [
            {"kind": "context", "field": "major", "text": "计算机科学与技术", "source_quote": "计算机科学与技术"},
            {"kind": "context", "field": "start_date", "text": "2023.09", "source_quote": "2023.09"}]}]})
        sparse = await extract_facts_patch(message, {}, model)
        facts = absorb_education(apply_fact_patch({}, sparse), {}, message)
        data = _serialize_to_resume_data({"collected_info": facts})
        data = overlay_education(facts, data)
        row = data["education"][0]
        self.assertEqual((row["school"], row["major"], row["degree"], row["start_date"], row["end_date"]),
                         ("示例甲大学", "计算机科学与技术", "本科", "2023.09", "2027.06"))
        self.assertIn("校级三等奖学金", row["highlights"])
        self.assertEqual(data["basic_info"]["availability"], "每周 6 天，6 个月起")
        self.assertEqual([link["text"] for link in data["basic_info"]["links"]], ["GitHub", "个人主页"])

    async def test_p2_reference_lists_do_not_become_candidate_facts(self):
        sparse = await extract_facts_patch("参考简历：\n" + LISTS, {}, ExtractModel({}))
        self.assertFalse(sparse["items"])
        self.assertFalse(deterministic_proposals({"human_text": LISTS, "source_kind": "reference"}))

    async def test_p2_bad_counter_unit_is_rejected_during_extraction(self):
        old = {"projects": [{"item_id": str(uuid.uuid4()), "name": "简历练习", "facts": [new_fact("做了简历练习")]}]}
        model = ExtractModel({"items": [{"section": "projects", "item_id": old["projects"][0]["item_id"],
              "name": "简历练习", "facts": [{"kind": "result_metric", "text": "累计10人使用、20次调用",
                                               "source_quote": "10个人 20次"}]}]})
        sparse = await extract_facts_patch("10个人 20次", old, model, previous_assistant="多少人用过，做过多少次简历任务？")
        self.assertFalse(sparse["items"])
        self.assertTrue(any("unsupported_quantity_unit" in str(entry) for entry in sparse["rejected"]))


class P2WriterTests(unittest.IsolatedAsyncioTestCase):
    async def test_p2_missing_second_pr_falls_back_to_its_own_facts(self):
        facts, _ = canonical_material()
        first = facts["open_source"][0]
        body = next(fact for fact in first["facts"] if not fact.get("field"))
        model = WriterModel({"match_plan": {}, "writer": {"items": [
            {"item_id": first["item_id"], "section": "open_source", "title": {"name": "自己", "role": "原型"},
             "bullets": [{"text": body["text"], "fact_ids": [body["id"]]}]}]}})
        result = await run_resume_writer(facts, {}, model)
        data = ResumeData.model_validate(result["resume_data"]).model_dump()
        self.assertEqual([row["name"] for row in data["open_source"]], ["StarBoat 项目", "CloudNote"])
        self.assertTrue(data["open_source"][1]["bullets"])
        self.assertEqual(len(data["awards"]), 3)
        self.assertEqual(data["education"][0]["major"], "计算机科学与技术")
        self.assertIn("奖学金", data["education"][0]["highlights"])
        self.assertEqual(data["basic_info"]["availability"], "每周 6 天，6 个月起")
        self.assertEqual(result["writer_diagnostic"]["status"], "model_completed")

    async def test_p2_writer_rejects_counter_reinterpretation_and_keeps_source(self):
        item_id = str(uuid.uuid4())
        fact = new_fact("10个人 20次", kind="result_metric")
        facts = {"projects": [{"item_id": item_id, "name": "简历练习", "facts": [fact]}]}
        model = WriterModel({"match_plan": {}, "writer": {"items": [
            {"item_id": item_id, "section": "projects", "bullets": [
                {"text": "累计10人使用、20次调用", "fact_ids": [fact["id"]]}]}]}})
        result = await run_resume_writer(facts, {}, model)
        body = "\n".join(result["resume_data"]["projects"][0]["polished_bullets"])
        self.assertIn("20次", body)
        self.assertNotIn("调用", body)
        self.assertTrue(result["guard_fallbacks"])

    async def test_p2_writer_keeps_deliberately_removed_new_sections_hidden(self):
        facts, data = canonical_material()
        data["sections"] = [{"id": "education", "kind": "education", "title": "教育背景"}]
        result = await run_resume_writer(facts, {}, WriterModel({"match_plan": {}, "writer": {"items": []}}),
                                         current_data=data)
        self.assertEqual([section["kind"] for section in result["resume_data"]["sections"]], ["education"])
        self.assertEqual(len(result["resume_data"]["open_source"]), 2)
        self.assertNotIn("StarBoat", render_resume("t001-jianyue", result["resume_data"]))

    async def test_p2_guard_does_not_need_model_to_detect_changed_units(self):
        fact = new_fact("10个人 20次", kind="result_metric")
        checked = await guard_bullets([{"text": "10人使用、20次请求", "fact_ids": [fact["id"]]}], [fact])
        self.assertEqual(checked[0]["verdict"], "fail")
        self.assertTrue(any("unsupported_quantity_unit" in value for value in checked[0]["lexical_errors"]))

    async def test_p2_section_advice_preserves_p1_read_only_intent(self):
        data = {"projects": [{"name": "StarBoat"}]}
        routed = await route_intents("你觉得这是开源贡献吗？", {"resume_data": data},
                                     WriterModel({"intents": ["edit"]}))
        self.assertIn("advice", routed["intents"])
        self.assertNotIn("edit", routed["intents"])
        self.assertFalse(routed["edit_plan"])

    async def test_p2_section_correction_routes_and_moves_without_writer(self):
        identifier = str(uuid.uuid4())
        quote = "StarBoat PR #5：修复文件校验，维护者已合并。"
        old = {"projects": [{"item_id": identifier, "name": "自己", "role": "原型", "polished_bullets": [quote]}]}
        facts = {"projects": [{"item_id": identifier, "name": "自己", "facts": [new_fact(quote)], "raw_notes": [quote]}]}
        state = {"stage": "EVALUATING", "resume_data": old, "collected_info": facts,
                 "messages": [HumanMessage(content="这是开源贡献，和项目经历一样是一个标题，不是自己")],
                 "resume_meta": {}, "resume_versions": []}
        with patch("app.agents.resume_graph._write_resume_update", AsyncMock(side_effect=AssertionError("writer called"))), \
             patch("app.agents.resume_graph._interview_suffix", return_value=("", [])), \
             patch("app.agents.tailor_pipeline.measure_rendered_pages", return_value={"status": "rendered", "pages": 1}):
            result = await evaluating_node(state, WriterModel({"intents": ["edit"], "edit_plan": []}))
        self.assertFalse(result["resume_data"]["projects"])
        moved = result["resume_data"]["open_source"][0]
        self.assertEqual((moved["item_id"], moved["name"], moved["pr"]), (identifier, "StarBoat", "PR #5"))
        self.assertEqual(moved["bullets"], [quote])
        self.assertEqual(result["collected_info"]["open_source"][0]["item_id"], identifier)
        self.assertTrue(result["changed"])
        self.assertEqual(result["writer_diagnostic"]["status"], "not_called")
        self.assertEqual(result["resume_versions"][0]["resume_data"], old)


class P2ProjectionTests(unittest.TestCase):
    def test_p2_education_overlay_keeps_legacy_description_honor(self):
        facts = {"education": [{"school": "示例甲大学", "degree": "本科", "major": "计算机科学与技术",
                 "start_date": "2023.09", "end_date": "2027.06", "facts": [new_fact("获得校级三等奖学金")]}]}
        data = overlay_education(facts, {"education": [{"school": "旧描述"}]})
        self.assertIn("校级三等奖学金", data["education"][0]["highlights"])
        self.assertEqual(data["education"][0]["start_date"], "2023.09")

    def test_p2_manual_edit_fields_and_links_reach_active_evidence(self):
        new = {"basic_info": {"availability": "每周5天", "links": [{"text": "主页", "url": "example.test"}]},
               "open_source": [{"name": "StarBoat", "pr": "PR #4", "summary": "7k★", "bullets": ["修复文件校验"]}],
               "awards": [{"name": "示例算法赛", "level": "省级三等奖", "date": "2026.04", "bullets": []}]}
        collected = append_editor_facts({}, {}, new, "synthetic-editor")
        active = active_facts_for_writer(collected)
        self.assertEqual(active["basic_info"], new["basic_info"])
        self.assertEqual({fact.get("field") for fact in active["awards"][0]["facts"]}, {"name", "level", "date"})
        canonical = _serialize_to_resume_data({"collected_info": collected})
        self.assertEqual(canonical["open_source"][0]["pr"], "PR #4")
        self.assertEqual(canonical["awards"][0]["level"], "省级三等奖")

    def test_p2_editor_keeps_two_prs_from_the_same_repository(self):
        data = {"open_source": [{"name": "StarBoat", "pr": "PR #1", "bullets": ["修复文件校验"]},
                                {"name": "StarBoat", "pr": "PR #2", "bullets": ["修复列表分页"]}]}
        collected = append_editor_facts({}, {}, data, "synthetic-editor")
        self.assertEqual(len(collected["open_source"]), 2)
        self.assertEqual(len({item["item_id"] for item in collected["open_source"]}), 2)
        self.assertEqual({item["pr"] for item in collected["open_source"]}, {"PR #1", "PR #2"})

    def test_p2_optional_schema_fields_load_old_data_without_migration(self):
        old = {"basic_info": {"name": "虚构小青"}, "projects": [{"name": "旧项目", "description": "旧内容"}]}
        data = ResumeData.model_validate(old)
        self.assertEqual(data.open_source, [])
        self.assertEqual(data.awards, [])
        dumped = data.model_dump()
        self.assertNotIn("links", dumped["basic_info"])
        self.assertNotIn("availability", dumped["basic_info"])
        self.assertNotIn("open_source", dumped)

    def test_p2_default_and_custom_section_visibility(self):
        _, data = canonical_material()
        self.assertIn("open_source", [section["kind"] for section in get_resume_sections(data)])
        data["sections"] = [{"id": "awards", "kind": "awards", "title": "获奖记录"}]
        text = docx_text(resume_data_to_docx(data))
        self.assertIn("获奖记录", text)
        self.assertNotIn("StarBoat", text)
        self.assertNotIn("示例甲大学", text)

    def test_p2_template_preview_and_word_keep_every_field(self):
        _, data = canonical_material()
        wanted = ["StarBoat", "PR #123", "7k★", "CloudNote", "PR #8", "Top100", "示例校算法赛", "二等奖",
                  "示例甲大学", "计算机科学与技术", "本科", "2023.09", "2027.06", "校级三等奖学金",
                  "https://github.com/fictional-example", "example.test", "每周 6 天，6 个月起"]
        for template in list_templates():
            html = render_resume(template["id"], data, interactive=True)
            for value in wanted:
                self.assertIn(value, html, (template["id"], value))
            self.assertIn('data-field-path="/open_source/1/name"', html)
            self.assertIn('data-field-path="/basic_info/links/1/url"', html)
        word = docx_text(resume_data_to_docx(ResumeData.model_validate(data).model_dump()))
        for value in wanted:
            self.assertIn(value, word)
            self.assertIn(value, resume_to_text(data))

    def test_p2_pdf_export_contains_the_same_sections(self):
        import pdfplumber
        _, data = canonical_material()
        with pdfplumber.open(io.BytesIO(render_resume_pdf("t001-jianyue", data))) as pdf:
            text = "\n".join(page.extract_text() or "" for page in pdf.pages)
        compact = "".join(text.split())
        for value in ("开源贡献", "竞赛荣誉", "CloudNote", "PR#8", "校级三等奖学金", "2023.09", "example.test", "每周6天"):
            self.assertIn(value, compact)

    def test_p2_legacy_word_template_uses_complete_export_for_new_columns(self):
        facts, _ = canonical_material()
        with tempfile.TemporaryDirectory() as directory:
            template = Path(directory) / "synthetic.docx"
            Document().save(template)
            text = docx_text(fill_template_bytes(template, facts))
        self.assertIn("CloudNote", text)
        self.assertIn("校级三等奖学金", text)
        self.assertIn("每周 6 天", text)

    def test_p2_education_only_legacy_word_keeps_dates_and_scholarship(self):
        facts = absorb_education({}, {}, "教育背景：" + EDUCATION)
        with tempfile.TemporaryDirectory() as directory:
            template = Path(directory) / "synthetic.docx"
            Document().save(template)
            text = docx_text(fill_template_bytes(template, facts))
        for value in ("示例甲大学", "计算机科学与技术", "2023.09", "2027.06", "校级三等奖学金"):
            self.assertIn(value, text)

    def test_p2_hidden_items_and_unsafe_links_are_not_rendered_as_active_content(self):
        _, data = canonical_material()
        data["open_source"][0]["hidden"] = True
        data["awards"][0]["hidden"] = True
        data["basic_info"]["links"] = [{"text": "<script>bad</script>", "url": "javascript:alert(1)"}]
        html = render_resume("t001-jianyue", data)
        self.assertNotIn("StarBoat", html)
        self.assertNotIn("Top100", html)
        self.assertNotIn('href="javascript:', html)
        self.assertIn("&lt;script&gt;", html)

    def test_p2_award_move_keeps_identity_body_and_produces_diff(self):
        identifier = str(uuid.uuid4())
        old = {"projects": [{"item_id": identifier, "name": "示例算法赛", "polished_bullets": ["省级三等奖"]}]}
        text = "示例算法赛这个放到竞赛荣誉"
        ops = choose_edit_plan(text, old, [])
        result = execute_edit_plan(old, {}, ops, instruction=text)
        self.assertEqual(result["data"]["awards"][0]["item_id"], identifier)
        self.assertEqual(result["data"]["awards"][0]["bullets"], ["省级三等奖"])
        self.assertTrue(document_diff(old, result["data"]))
        self.assertEqual(result["writer_requests"], [])

    def test_p2_move_recovers_repository_from_a_multiline_original_note(self):
        identifier = str(uuid.uuid4())
        old = {"projects": [{"item_id": identifier, "name": "自己", "role": "原型", "polished_bullets": ["修复文件校验"]}]}
        facts = {"projects": [{"item_id": identifier, "name": "自己", "facts": [new_fact("修复文件校验")],
                 "raw_notes": ["开源贡献：\n1. StarBoat（7k★）https://github.com/example/StarBoat/pull/4：修复文件校验。\n技能：PRD、原型"]}]}
        instruction = "这是开源贡献，不是自己"
        result = execute_edit_plan(old, {}, choose_edit_plan(instruction, old, []),
                                   collected_info=facts, instruction=instruction)
        row = result["data"]["open_source"][0]
        self.assertEqual((row["name"], row["summary"], row["pr"]),
                         ("StarBoat", "7k★", "https://github.com/example/StarBoat/pull/4"))

    def test_p2_ambiguous_or_unrequested_move_does_not_touch_data(self):
        old = {"projects": [{"name": "甲项目"}, {"name": "乙项目"}]}
        ops = choose_edit_plan("这个放到竞赛荣誉", old, [])
        result = execute_edit_plan(old, {}, ops, instruction="这个放到竞赛荣誉")
        self.assertEqual(result["data"], old)
        self.assertEqual(result["operation_results"][0]["status"], "needs_confirmation")
        result = execute_edit_plan(old, {}, [{"op": "move_item", "target": "甲项目", "args": {"section": "open_source"}}], instruction="精简甲项目")
        self.assertEqual(result["data"], old)

    def test_p2_moving_one_item_keeps_a_separate_rewrite_request(self):
        old = {"projects": [{"name": "StarBoat"}, {"name": "课程工具"}]}
        instruction = "StarBoat 放到开源贡献；精简课程工具"
        ops = choose_edit_plan(instruction, old, [{"op": "compress", "target": "课程工具", "args": {}}])
        self.assertTrue(any(op["op"] == "compress" and op["target"] == "课程工具" for op in ops))
        self.assertTrue(any(op["op"] == "move_item" for op in ops))

    def test_p2_changed_unit_detection_allows_original_units_and_chinese_digits(self):
        for text in ("10人使用、20次调用", "10人使用、20次请求", "10条、20次", "10人、20个"):
            self.assertTrue(any("quantity_unit" in value for value in lexical_fact_errors(text, "10个人 20次")), text)
        for text, quote in (("10人，20次", "10个人 20次"), ("10人，20次", "十个人，二十次"),
                            ("调用20次", "20次调用"), ("20次请求", "请求了20次")):
            self.assertEqual(lexical_fact_errors(text, quote), [], (text, quote))

    def test_p2_workspace_and_bullet_metadata_include_contributions(self):
        facts, data = canonical_material()
        view = build_workspace({"collected_info": facts}, data)
        self.assertEqual(len(view["experiences"]["items"]), 5)
        self.assertEqual(view["experiences"]["items"][1]["name"], "CloudNote")
        row = data["open_source"][0]
        from app.services.resume_versions import text_hash
        meta = {"bullets": {row["item_id"]: [{"text_hash": text_hash(row["bullets"][0]), "fact_ids": ["synthetic"]}]}}
        self.assertIn(row["item_id"], valid_bullet_metadata(data, meta)["bullets"])


if __name__ == "__main__":
    unittest.main()
