"""Hidden backup rows stay in ResumeData but never enter candidate output."""

import io
from pathlib import Path
import re
import shutil
import subprocess
import sys
import unittest

from docx import Document

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "backend"))

from app.schemas.editor import ResumeData  # noqa: E402
from app.services.renderer import list_templates, render_resume  # noqa: E402
from app.services.resume_pdf import render_resume_pdf  # noqa: E402
from app.services.resume_review import resume_to_text, review_resume  # noqa: E402
from app.tools.resume_data_docx import resume_data_to_docx  # noqa: E402


def sample() -> dict:
    return ResumeData.model_validate({
        "basic_info": {"name": "Test Person"},
        "education": [
            {"item_id": "00000000-0000-0000-0000-000000000001",
             "school": "HIDDEN_SCHOOL_A", "degree": "HIDDEN_DEGREE_A", "hidden": True,
             "links": [{"text": "Source", "url": "https://example.test/hidden"}]},
            {"item_id": "00000000-0000-0000-0000-000000000002",
             "school": "VISIBLE_SCHOOL_B", "degree": "VISIBLE_DEGREE_B"},
        ],
        "work_experience": [
            {"item_id": "00000000-0000-0000-0000-000000000003",
             "company": "HIDDEN_WORK_A", "bullets": ["HIDDEN_WORK_BULLET"], "hidden": True},
            {"item_id": "00000000-0000-0000-0000-000000000004",
             "company": "VISIBLE_WORK_B", "bullets": ["VISIBLE_WORK_BULLET"]},
        ],
        "projects": [
            {"item_id": "00000000-0000-0000-0000-000000000005",
             "name": "HIDDEN_PROJECT_A", "description": "HIDDEN_PROJECT_BODY", "hidden": True},
            {"item_id": "00000000-0000-0000-0000-000000000006",
             "name": "VISIBLE_PROJECT_B", "description": "VISIBLE_PROJECT_BODY",
             "links": [{"text": "Project", "url": "https://example.test/project"}]},
        ],
    }).model_dump()


HIDDEN_MARKERS = (
    "HIDDEN_SCHOOL_A", "HIDDEN_DEGREE_A", "HIDDEN_WORK_A",
    "HIDDEN_WORK_BULLET", "HIDDEN_PROJECT_A", "HIDDEN_PROJECT_BODY",
)


class HiddenResumeItemsTest(unittest.TestCase):
    def test_schema_keeps_hidden_rows_and_structured_links(self):
        data = sample()
        self.assertEqual(len(data["projects"]), 2)
        self.assertTrue(data["projects"][0]["hidden"])
        self.assertEqual(data["projects"][1]["links"],
                         [{"text": "Project", "url": "https://example.test/project"}])
        self.assertEqual(data["education"][0]["links"][0]["url"], "https://example.test/hidden")
        self.assertEqual(data["projects"][0]["item_id"], "00000000-0000-0000-0000-000000000005")

    def test_all_html_templates_skip_hidden_and_preserve_source_paths(self):
        data = sample()
        self.assertIn("t003-shenlan", {template["id"] for template in list_templates()})
        for template in list_templates():
            with self.subTest(template=template["id"]):
                html = render_resume(template["id"], data, interactive=True)
                for marker in HIDDEN_MARKERS:
                    self.assertNotIn(marker, html)
                for collection, field in (("education", "school"), ("education", "degree"),
                                          ("work_experience", "company"), ("projects", "name")):
                    self.assertIn(f'data-field-path="/{collection}/1/{field}"', html)
                    self.assertNotIn(f'data-field-path="/{collection}/0/{field}"', html)
                self.assertIn("VISIBLE_SCHOOL_B", html)
                self.assertIn("VISIBLE_PROJECT_B", html)
        self.assertEqual(len(data["projects"]), 2)
        self.assertTrue(data["projects"][0]["hidden"])
        self.assertEqual(data["projects"][0]["item_id"], "00000000-0000-0000-0000-000000000005")

    def test_all_hidden_rows_remove_builtin_heading_and_education_header_aliases(self):
        data = sample()
        for kind in ("education", "work_experience", "projects"):
            for item in data[kind]:
                item["hidden"] = True
        for template in list_templates():
            with self.subTest(template=template["id"]):
                html = render_resume(template["id"], data)
                self.assertNotIn("教育背景", html)
                self.assertNotIn("实习经历", html)
                self.assertNotIn("项目经历", html)
                self.assertIsNone(re.search(r'data-field-path="/(?:education|work_experience|projects)/', html))
                self.assertNotIn("VISIBLE_SCHOOL_B", html)
                self.assertNotIn("VISIBLE_DEGREE_B", html)

    def test_word_and_plain_text_skip_hidden_rows_and_empty_headings(self):
        data = sample()
        plain = resume_to_text(data)
        word = "\n".join(paragraph.text for paragraph in Document(io.BytesIO(resume_data_to_docx(data))).paragraphs)
        for output in (plain, word):
            for marker in HIDDEN_MARKERS:
                self.assertNotIn(marker, output)
            self.assertIn("VISIBLE_SCHOOL_B", output)
            self.assertIn("VISIBLE_WORK_BULLET", output)
            self.assertIn("VISIBLE_PROJECT_BODY", output)
        for kind in ("education", "work_experience", "projects"):
            for item in data[kind]:
                item["hidden"] = True
        plain = resume_to_text(data)
        word = "\n".join(paragraph.text for paragraph in Document(io.BytesIO(resume_data_to_docx(data))).paragraphs)
        for output in (plain, word):
            for heading in ("教育背景", "实习经历", "项目经历"):
                self.assertNotIn(heading, output)
        self.assertIn("补充至少一段", "\n".join(review_resume(data)))

    @unittest.skipUnless(shutil.which("pdftotext"), "Poppler pdftotext not installed")
    def test_production_pdf_for_all_six_templates_has_no_hidden_text(self):
        data = sample()
        for template in list_templates():
            with self.subTest(template=template["id"]):
                pdf = render_resume_pdf(template["id"], data)
                output = subprocess.run(["pdftotext", "-", "-"], input=pdf, capture_output=True,
                                        check=True, timeout=30).stdout.decode("utf-8")
                for marker in HIDDEN_MARKERS:
                    self.assertNotIn(marker, output)
                self.assertIn("VISIBLE_PROJECT_B", output)


if __name__ == "__main__":
    unittest.main()
