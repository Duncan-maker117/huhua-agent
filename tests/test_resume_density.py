"""One-page layout selection with fictional content and real export geometry."""

from copy import deepcopy
import io
from pathlib import Path
import re
import sys
import tempfile
import unittest
from unittest.mock import AsyncMock, patch
import zipfile

from docx import Document
from langchain_core.messages import HumanMessage
from lxml.html import fromstring
import pdfplumber

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / 'backend'))

from app.agents import resume_graph as graph
from app.agents.router import route_intents
from app.agents.tailor_pipeline import run_tailor_pipeline
from app.agents.turn_actions import deterministic_compound_plan
from app.schemas.editor import ResumeData
from app.services.edit_plan import execute_edit_plan
from app.services.renderer import list_templates, render_resume
from app.services.resume_capacity import measure_rendered_pages, plan_compression, revalidate_committed_page_check, resume_data_hash
from app.services.resume_layout import density_operations, resume_density
from app.services.resume_pdf import count_resume_pdf_pages, render_resume_pdf
from app.services.resume_versions import restore_payload
from app.tools.docx_template_filler import fill_template_bytes
from app.tools.resume_data_docx import resume_data_to_docx
import test_product_security_flow as product_fixture


def fictional_resume(projects=5):
    return {
        'basic_info': {'name': '虚构候选人林栩', 'objective': '应用开发实习生',
                       'phone': '10000000000', 'email': 'lin@example.test',
                       'links': [{'text': '作品', 'url': 'https://example.test/portfolio'}]},
        'education': [{'school': '虚构理工大学', 'major': '计算机科学', 'degree': '本科',
                       'start_date': '2023.09', 'end_date': '2027.06'}],
        'projects': [{'name': f'虚构检索项目{i}', 'item_id': str(i), 'role': '开发', 'polished_bullets': [
            f'样例{j}：整理虚构测试样例并归类反馈，使用 Python 实现检索接口，记录实验结果及功能限制。' * 2
            for j in range(3)]} for i in range(projects)],
        'open_source': [{'name': 'fictional-search', 'summary': '虚构检索工具', 'pr': 'PR #42',
                         'bullets': ['修复虚构接口的分页条件并补充测试。'],
                         'links': [{'text': 'PR #42', 'url': 'https://example.test/fictional-search/pull/42'}]}],
        'skills': ['Python / SQL / Git'],
    }


class NoModel:
    def __init__(self):
        self.calls = 0

    def with_structured_output(self, *args, **kwargs):
        self.calls += 1
        raise AssertionError('Layout editing must never request a model')

    async def ainvoke(self, *args, **kwargs):
        self.calls += 1
        raise AssertionError('Layout editing must never request a model')


class DensityLayoutTests(unittest.TestCase):
    def test_actual_two_pages_selects_first_feasible_density_without_touching_content(self):
        for projects, expected in ((4, 1), (5, 2), (6, 3)):
            data = fictional_resume(projects)
            original = deepcopy(data)
            counts = [count_resume_pdf_pages('t001-jianyue', {**data, 'layout': {'density': i}})
                      for i in range(4)]
            self.assertEqual(counts[0], 2)
            self.assertEqual(counts.index(1), expected)
            measured = measure_rendered_pages(data, {'page_limit': 1})
            self.assertEqual(measured['pages'], 1)
            self.assertEqual(data['layout'], {'density': expected, 'density_source': 'auto'})
            self.assertEqual(measured['density_attempts'], [{'density': i, 'pages': counts[i]}
                                                           for i in range(expected + 1)])
            self.assertEqual(measured['data_hash'], resume_data_hash(data))
            self.assertEqual({k: v for k, v in data.items() if k != 'layout'}, original)

    def test_short_content_keeps_default_and_automatic_density_can_relax_again(self):
        data = fictional_resume(1)
        measured = measure_rendered_pages(data, {'page_limit': 1})
        self.assertEqual(measured['density'], 0)
        self.assertNotIn('layout', data)
        data['layout'] = {'density': 3, 'density_source': 'auto'}
        self.assertEqual(measure_rendered_pages(data, {'page_limit': 1})['density'], 0)
        self.assertEqual(data['layout']['density'], 0)

    def test_user_density_is_locked_even_when_it_overflows_or_could_relax(self):
        for projects, density, pages in ((5, 0, 2), (1, 3, 1)):
            data = fictional_resume(projects)
            data['layout'] = {'density': density, 'density_source': 'user'}
            original = deepcopy(data)
            measured = measure_rendered_pages(data, {'page_limit': 1})
            self.assertEqual(measured['pages'], pages)
            self.assertEqual(measured['density_attempts'], [{'density': density, 'pages': pages}])
            self.assertEqual(data, original)

    def test_two_page_constraint_does_not_compact(self):
        data = fictional_resume(5)
        result = measure_rendered_pages(data, {'page_limit': 2})
        self.assertEqual(result['pages'], 2)
        self.assertEqual(result['density'], 0)
        self.assertNotIn('layout', data)

    def test_pdf_bytes_can_trigger_layout_retry_and_hash_matches_final_draft(self):
        data = fictional_resume(5)
        pdf = render_resume_pdf('t001-jianyue', data)
        measured = measure_rendered_pages(data, {'page_limit': 1}, pdf_bytes=pdf)
        self.assertEqual(measured['density'], 2)
        self.assertEqual(measured['pages'], 1)
        self.assertEqual(measured['data_hash'], resume_data_hash(data))
        self.assertEqual(measured['source'], 'production_render_resume_pdf')

    def test_failed_retry_keeps_existing_layout_and_never_claims_fit(self):
        data = fictional_resume(5)
        original = deepcopy(data)
        from app.services.resume_pdf import measure_resume_pdf_layout
        def fail(template, candidate, pages):
            if resume_density(candidate) == 1:
                raise RuntimeError('layout unavailable')
            return measure_resume_pdf_layout(template, candidate, pages)
        with patch('app.services.resume_pdf.measure_resume_pdf_layout', side_effect=fail):
            result = measure_rendered_pages(data, {'page_limit': 1})
        self.assertEqual(result['status'], 'render_error')
        self.assertIsNone(result['within_page_limit'])
        self.assertEqual(data, original)

    def test_legacy_cached_overflow_is_retried_on_save_even_if_content_is_unchanged(self):
        data = fictional_resume(5)
        constraints = {'page_limit': 1, 'template_id': 't001-jianyue'}
        old = {'page_check': {'status': 'rendered', 'pages': 2, 'within_page_limit': False,
                             **constraints, 'data_hash': resume_data_hash(data)}}
        result = revalidate_committed_page_check(data, old, constraints)
        self.assertEqual(result['page_check']['pages'], 1)
        self.assertEqual(data['layout']['density'], 2)
        self.assertEqual(result['page_check']['data_hash'], resume_data_hash(data))

    def test_all_registered_templates_share_density_variables_and_body_floor(self):
        for template in list_templates():
            for density in range(4):
                data = fictional_resume(1)
                data['layout'] = {'density': density}
                html = render_resume(template['id'], data)
                root = fromstring(html)
                self.assertEqual(root.find('body').get('data-density'), str(density))
                variables = dict(re.findall(r'--([\w-]+):\s*([^;]+);', html))
                for name in ('body-font-size', 'content-font-size', 'item-title-size', 'summary-font-size'):
                    self.assertGreaterEqual(float(variables[name][:-2]), 9.5)
                for name in ('line-height', 'item-gap', 'li-gap', 'section-bottom', 'section-padding-y', 'page-x'):
                    self.assertIn(f'var(--{name})', html)
                if density == 0:
                    self.assertEqual(html, render_resume(template['id'], {k: v for k, v in data.items() if k != 'layout'}))
                if density == 3:
                    self.assertEqual(variables['body-font-size'], '9.5pt')

    def test_contribution_links_are_inline_and_header_urls_stay_whole(self):
        data = fictional_resume(1)
        data['basic_info']['links'].append({'text': '长链接', 'url': 'https://example.test/' + 'long-' * 60})
        for template in list_templates():
            root = fromstring(render_resume(template['id'], data, interactive=True))
            name = root.xpath('.//*[@data-field-path="/open_source/0/name"]')[0]
            self.assertEqual(name.getparent().tag, 'a')
            self.assertEqual(name.getparent().get('href'), data['open_source'][0]['links'][0]['url'])
            url = root.xpath('.//*[@data-field-path="/open_source/0/links/0/url"]')[0]
            self.assertFalse(url.xpath('ancestor::*[contains(@class, "item-sub")]'))
            self.assertIn('white-space: nowrap', root.text_content())
            long = root.xpath('.//*[@data-field-path="/basic_info/links/1/url"]/a')[0]
            self.assertEqual(long.text_content(), '链接')
            self.assertEqual(long.get('href'), data['basic_info']['links'][1]['url'])

    def test_generic_and_filled_word_read_density_and_keep_full_hyperlinks(self):
        for density in range(4):
            data = fictional_resume(1)
            data['layout'] = {'density': density, 'density_source': 'user'}
            generic = Document(io.BytesIO(resume_data_to_docx(data)))
            self.assertEqual(generic.core_properties.keywords, f'resume-density={density}')
            self.assertGreaterEqual(generic.styles['Normal'].font.size.pt, 9.5)
            if density:
                expected = (1.1, 1.05, 1.0)[density - 1]
                self.assertAlmostEqual(generic.styles['Normal'].paragraph_format.line_spacing, expected, delta=.003)
            with zipfile.ZipFile(io.BytesIO(resume_data_to_docx(data))) as archive:
                self.assertIn(data['open_source'][0]['links'][0]['url'], archive.read('word/_rels/document.xml.rels').decode())
            with tempfile.TemporaryDirectory() as directory:
                path = Path(directory) / 'fictional.docx'
                template = Document()
                template.add_paragraph('虚构模版正文')
                template.save(path)
                filled = Document(io.BytesIO(fill_template_bytes(path, {'layout': data['layout']})))
                if density:
                    self.assertEqual(filled.core_properties.keywords, f'resume-density={density}')
                    self.assertGreaterEqual(filled.styles['Normal'].font.size.pt, 9.5)
                    self.assertAlmostEqual(filled.paragraphs[0].paragraph_format.line_spacing, expected, delta=.003)
                fallback = Document(io.BytesIO(fill_template_bytes(path, data)))
                self.assertEqual(fallback.core_properties.keywords, f'resume-density={density}')

    def test_schema_preserves_optional_density_and_rejects_invalid_values(self):
        self.assertNotIn('layout', ResumeData().model_dump())
        for density in range(4):
            data = {'layout': {'density': density, 'density_source': 'user'}}
            self.assertEqual(ResumeData.model_validate(data).model_dump()['layout'], data['layout'])
        for density in (-1, 4, True, '2'):
            with self.assertRaises(ValueError):
                ResumeData.model_validate({'layout': {'density': density}})


class DensityEditTests(unittest.IsolatedAsyncioTestCase):
    async def test_literal_density_edits_use_existing_executor_and_never_call_model(self):
        for text, initial, expected in (('字大一点', 2, 1), ('字小一点', 1, 2),
                                        ('排紧凑点', 2, 3), ('排松一点', 2, 1),
                                        ('字大一点', 0, 0), ('字小一点', 3, 3)):
            data = fictional_resume(1)
            data['layout'] = {'density': initial, 'density_source': 'auto'}
            model = NoModel()
            routed = await route_intents(text, {'resume_data': data}, model)
            self.assertEqual(routed['edit_plan'], deterministic_compound_plan(text, data))
            executed = execute_edit_plan(data, {}, routed['edit_plan'], instruction=text)
            self.assertEqual(executed['data']['layout'], {'density': expected, 'density_source': 'user'})
            self.assertEqual(executed['writer_requests'], [])
            state = {**graph._init_state(), 'stage': 'EVALUATING', 'resume_data': data,
                     'messages': [HumanMessage(content=text)]}
            with patch('app.agents.tailor_pipeline.run_resume_writer', AsyncMock(side_effect=AssertionError('writer called'))):
                updated = await graph.evaluating_node(state, model)
            self.assertEqual(updated['resume_data']['layout'], {'density': expected, 'density_source': 'user'})
            self.assertNotIn('当前排版', updated['messages'][-1].content)
            self.assertNotIn('实际为 1 页', updated['messages'][-1].content)
            self.assertEqual(model.calls, 0)
            self.assertEqual(updated['resume_data']['projects'], data['projects'])
            self.assertTrue(updated['changed'])
            restored = restore_payload({**updated, 'collected_info': {}})
            self.assertEqual(restored['resume_data']['layout'], data['layout'])

    async def test_auto_request_unlocks_and_reselects_without_model(self):
        data = fictional_resume(5)
        data['layout'] = {'density': 0, 'density_source': 'user'}
        model = NoModel()
        state = {**graph._init_state(), 'stage': 'EVALUATING', 'resume_data': data,
                 'messages': [HumanMessage(content='自动调整')]}
        with patch('app.agents.tailor_pipeline.run_resume_writer', AsyncMock(side_effect=AssertionError('writer called'))):
            updated = await graph.evaluating_node(state, model)
        self.assertEqual(updated['resume_data']['layout'], {'density': 2, 'density_source': 'auto'})
        self.assertEqual(updated['page_check']['data_hash'], resume_data_hash(updated['resume_data']))
        self.assertEqual(model.calls, 0)

    async def test_unrequested_and_negated_density_edits_are_rejected(self):
        for text in ('不要字小一点', '你觉得字大一点合适吗', '比如排紧凑点'):
            self.assertEqual(density_operations(text), [])
        data = fictional_resume(1)
        result = execute_edit_plan(data, {}, [{'op': 'set_constraint', 'target': 'density', 'args': {'density_delta': 1}}],
                                   instruction='修改项目名称')
        self.assertEqual(result['data'], data)
        self.assertEqual(result['operation_results'][0]['status'], 'cannot')

    async def test_pipeline_compresses_only_after_density_three_still_overflows(self):
        for projects in (5, 10):
            data = fictional_resume(projects)
            calls = []
            async def fake_writer(*args, **kwargs):
                calls.append(kwargs['scope'])
                return {'resume_data': deepcopy(kwargs['current_data']), 'resume_meta': {},
                        'match_plan': {}, 'writer_diagnostic': {'status': 'fake_completed'}}
            def compression(candidate, *args, **kwargs):
                self.assertEqual(resume_density(candidate), 3)
                return plan_compression(candidate, *args, **kwargs)
            with patch('app.agents.tailor_pipeline.run_resume_writer', side_effect=fake_writer), \
                    patch('app.agents.tailor_pipeline.plan_compression', side_effect=compression) as compress:
                result = await run_tailor_pipeline({'resume_data': data}, data, NoModel(), instruction='生成一页')
            if projects == 5:
                compress.assert_not_called()
                self.assertEqual(result['page_check']['pages'], 1)
            else:
                self.assertTrue(compress.called)
                self.assertGreater(result['page_check']['pages'], 1)
                self.assertEqual(result['resume_data']['layout']['density'], 3)
                self.assertTrue(result['compression_retention_checks'])
                self.assertIn('由你指定移除', graph._writer_reply(result))
            self.assertEqual(calls, ['whole'])


class DensityApiTests(unittest.IsolatedAsyncioTestCase):
    asyncSetUp = product_fixture.ProductSecurityFlow.asyncSetUp
    cleanup = product_fixture.ProductSecurityFlow.cleanup
    login = product_fixture.ProductSecurityFlow.login
    fixture = product_fixture.ProductSecurityFlow.fixture

    async def test_saved_preview_pdf_and_word_share_manual_density(self):
        sid, _ = await self.fixture(marker='FICTIONAL_DENSITY')
        current = (await self.a.get(f'/api/resume/current?session_id={sid}')).json()
        data = deepcopy(current['data'])
        data['layout'] = {'density': 3, 'density_source': 'user'}
        saved = await self.a.put('/api/resume/current', json={
            'session_id': sid, 'expected_revision': current['revision'], 'resume_data': data})
        self.assertEqual(saved.status_code, 200, saved.text)
        self.assertEqual(saved.json()['data']['layout'], data['layout'])
        self.assertEqual(saved.json()['page_check']['density'], 3)
        # Exports and the saved preview must use the committed density even if
        # an old client submits another value with the download request.
        data['layout']['density'] = 0
        preview = await self.a.post(f'/api/resume/render/t001-jianyue?session_id={sid}', json=data)
        pdf = await self.a.post(f'/api/resume/export/pdf/t001-jianyue?session_id={sid}', json=data)
        word = await self.a.post(f'/api/resume/export/docx/t001-jianyue?session_id={sid}', json=data)
        self.assertEqual(preview.status_code, 200)
        self.assertIn('data-density="3"', preview.text)
        self.assertIn('--body-font-size: 9.5pt;', preview.text)
        self.assertEqual(pdf.status_code, 200)
        with pdfplumber.open(io.BytesIO(pdf.content)) as document:
            self.assertEqual(len(document.pages), saved.json()['page_check']['pages'])
            self.assertTrue(any(abs(char['size'] - 9.5) < .01 for char in document.pages[0].chars))
        self.assertEqual(word.status_code, 200)
        doc = Document(io.BytesIO(word.content))
        self.assertEqual(doc.core_properties.keywords, 'resume-density=3')
        self.assertEqual(doc.styles['Normal'].font.size.pt, 9.5)
        self.assertEqual(doc.styles['List Bullet'].font.size.pt, 9.5)
        # A legacy editor that has no layout field must not erase a user lock.
        data = deepcopy(saved.json()['data'])
        data.pop('layout')
        again = await self.a.put('/api/resume/current', json={
            'session_id': sid, 'expected_revision': saved.json()['revision'], 'resume_data': data})
        self.assertEqual(again.status_code, 200, again.text)
        self.assertEqual(again.json()['data']['layout'], {'density': 3, 'density_source': 'user'})

    async def test_editor_save_selects_layout_and_binds_checkpoint_and_sql_hash(self):
        sid, _ = await self.fixture(marker='FICTIONAL_AUTO_LAYOUT')
        current = (await self.a.get(f'/api/resume/current?session_id={sid}')).json()
        data = fictional_resume(5)
        for item in data['projects']:
            item.pop('item_id')
        removed = [item['item_id'] for item in current['data']['projects'] if item.get('item_id')]
        saved = await self.a.put('/api/resume/current', json={
            'session_id': sid, 'expected_revision': current['revision'], 'resume_data': data,
            'removed_item_ids': removed})
        self.assertEqual(saved.status_code, 200, saved.text)
        committed = saved.json()['data']
        self.assertEqual(committed['layout'], {'density': 2, 'density_source': 'auto'})
        self.assertEqual(saved.json()['page_check']['pages'], 1)
        self.assertEqual(saved.json()['page_check']['data_hash'], resume_data_hash(committed))
        snapshot = await self.graph.aget_state({'configurable': {'thread_id': sid}})
        self.assertEqual(snapshot.values['resume_data'], committed)
        reloaded = (await self.a.get(f'/api/resume/current?session_id={sid}')).json()
        self.assertEqual(reloaded['data'], committed)


if __name__ == '__main__':
    unittest.main()
