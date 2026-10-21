"""Technology lines: offline source, writer, editing, template and export regressions."""
from copy import deepcopy
import io
import json
from pathlib import Path
import shutil
import subprocess
import sys
from types import SimpleNamespace
import unittest
from unittest.mock import AsyncMock, patch
from uuid import uuid4

ROOT = Path(__file__).resolve().parents[1]
sys.path[:0] = [str(ROOT / 'backend'), str(ROOT / 'tests')]
from docx import Document  # noqa: E402
from langchain_core.messages import HumanMessage  # noqa: E402
from app.agents import orchestrator_tools as tools, resume_graph as graph  # noqa: E402
from app.agents.resume_writer import run_resume_writer  # noqa: E402
from app.mcp.server import MCPResume  # noqa: E402
from app.mcp.writing_guide import writing_guide  # noqa: E402
from app.schemas.editor import ProjectItem, ResumeData, WorkItem  # noqa: E402
from app.services.fact_guard import guard_bullets  # noqa: E402
from app.services.fact_coverage import ensure_fact_coverage  # noqa: E402
from app.services.mcp_service import MCPService, resume_claims  # noqa: E402
from app.services.resume_facts import append_editor_facts, normalize_fact_store  # noqa: E402
from app.services.resume_versions import document_diff, prepare_version, restore_payload, text_hash, valid_bullet_metadata  # noqa: E402
from app.services.tech_stack import (  # noqa: E402
    apply_stack_edit, extract_stack, normalize_stack, parse_stack_edit,
                                     project_item_stack, pure_stack, supporting_ids)
from app.services.renderer import TEMPLATE_DIR, render_resume  # noqa: E402
from app.tools.resume_data_docx import resume_data_to_docx  # noqa: E402
from test_fact_guard import decision  # noqa: E402
from test_orchestrator_tools import measured  # noqa: E402

P, W = str(uuid4()), str(uuid4())


def fact(fid, text, **extra):
    return {'id': fid, 'text': text, 'source_quote': text, 'status': 'active', 'kind': 'method_tool', **extra}


def collected():
    return {'projects': [{'item_id': P, 'name': '胡话简历', 'facts': [
        fact('f1', '这个项目使用了 mcp、toolcalling、longchain、长期记忆'),
        fact('f2', '实现简历导出并整理核对记录', kind='action')]}],
        'work_experience': [{'item_id': W, 'company': '合成公司', 'title': '实习',
            'facts': [fact('f3', '技术栈：Python、FastAPI'), fact('f4', '整理接口文档并交付核对表', kind='action')]}]}


def document():
    return ResumeData.model_validate({'projects': [{'item_id': P, 'name': '胡话简历',
        'polished_bullets': ['实现简历导出并整理核对记录']}], 'work_experience': [
        {'item_id': W, 'company': '合成公司', 'title': '实习', 'bullets': ['整理接口文档并交付核对表']}]}).model_dump()


class WriterStub:
    def __init__(self, proposal=True):
        self.proposal = proposal
        self.calls = []

    def with_structured_output(self, schema):
        async def invoke(messages):
            payload = json.loads(messages[-1].content)
            self.calls.append(schema.__name__)
            if schema.__name__ == 'StructurePlan':
                return {'recommended': {'bullet_style': 'plain', 'section_order': ['projects', 'work_experience'],
                    'experience_plan': [{'item_id': iid, 'treatment': 'keep', 'target_bullets': 2} for iid in (P, W)]}}
            if schema.__name__ in {'OnePassOutput', 'WriterDraft'}:
                items = []
                for section in ('projects', 'work_experience'):
                    for row in payload['evidence'].get(section) or []:
                        refs = [f['id'] for f in row['facts']]
                        items.append({'item_id': row['item_id'], 'section': section,
                            'tech_stack': ['longchain', 'InventedCloud'] if self.proposal and section == 'projects' else [],
                            'tech_stack_fact_ids': refs,
                            'bullets': [{'text': f['text'], 'fact_ids': [f['id']]} for f in row['facts']]})
                writer = {'items': items}
                return {'match_plan': {'bullet_style': 'plain'}, 'writer': writer} if schema.__name__ == 'OnePassOutput' else writer
            if schema.__name__ == 'FactSlots':
                from app.services.labeled_writer import rule_slot
                return {'assignments': [{'fact_id': f['id'], 'slot': rule_slot(f) or '决策'} for f in payload['facts']]}
            if schema.__name__ == 'SemanticBatch':
                return {'decisions': [decision(row['id']) for row in payload]}
            if schema.__name__ == 'RepairOutput':
                return {'bullets': []}
            raise AssertionError(schema.__name__)
        return SimpleNamespace(ainvoke=invoke)


class TechSchemaTests(unittest.TestCase):
    def test_old_and_empty_omit_field(self):
        for schema in (ProjectItem, WorkItem):
            self.assertNotIn('tech_stack', schema().model_dump())
            self.assertNotIn('tech_stack', schema(tech_stack=[]).model_dump())
        self.assertNotIn('tech_stack', ResumeData(projects=[{}]).model_dump()['projects'][0])

    def test_aliases_order_dedup_cap_and_unknown(self):
        self.assertEqual(normalize_stack(['mcp', ' toolcalling ', 'MCP', 'longchain', 'LangChain', '长期记忆', 'NebulaDB']),
                         ['MCP', 'Tool Calling', 'LangChain', '长期记忆', 'NebulaDB'])
        self.assertEqual(len(ProjectItem(tech_stack=['T' + str(i) for i in range(12)]).tech_stack), 8)
        self.assertEqual(normalize_stack(['记忆']), ['记忆'])  # No invented long-term capability.

    def test_mcp_exposes_both_fields(self):
        schema = MCPResume.model_json_schema()['$defs']
        for cls in ('ProjectItem', 'WorkItem'):
            self.assertEqual(schema[cls]['properties']['tech_stack']['type'], 'array')
        self.assertEqual(MCPResume(projects=[{'tech_stack': ['longchain']}]).projects[0].tech_stack, ['LangChain'])

    def test_only_lists_migrate(self):
        self.assertEqual(pure_stack('技术栈：mcp、longchain、MCP'), ['MCP', 'LangChain'])
        self.assertEqual(pure_stack('技术：基于 MCP 实现调用'), [])
        self.assertEqual(pure_stack('MCP、Tool Calling', bare=True), ['MCP', 'Tool Calling'])

    def test_extract_original_quote_and_mentions_order(self):
        self.assertEqual(extract_stack(collected()['projects'][0]['facts'][0]),
                         ['MCP', 'Tool Calling', 'LangChain', '长期记忆'])
        self.assertEqual(extract_stack(fact('f', '技术栈：NebulaDB、ZetaSDK')), ['NebulaDB', 'ZetaSDK'])
        self.assertEqual(extract_stack(fact('f', '技术栈：MCP', source_quote='没有用 MCP')), [])

    def test_negative_planned_partial_names_do_not_support(self):
        for quote in ['没用 MCP', '计划使用 MCP、LangChain', '不使用 MCP', '不是 MCP', '例如 MCP', '用 MCProxy']:
            self.assertFalse(supporting_ids('MCP', [fact('f', quote)]), quote)
        self.assertFalse(supporting_ids('Java', [fact('f', '使用 JavaScript')]))

    def test_metadata_invalidated_on_manual_rewrite(self):
        data = document()
        data['projects'][0]['tech_stack'] = ['MCP', 'Python']
        meta = {'tech_stack': {P: [{'text_hash': text_hash('MCP'), 'fact_ids': ['f1']},
                                 {'text_hash': text_hash('LangChain'), 'fact_ids': ['f1']}]}}
        self.assertEqual(valid_bullet_metadata(data, meta)['tech_stack'][P][1], {})

    def test_editor_input_and_legacy_source_are_recorded(self):
        data, old = document(), document()
        data['projects'][0]['tech_stack'] = ['Rust']
        store = append_editor_facts(collected(), old, data, 'editor-test')
        entry = next(f for f in store['projects'][0]['facts'] if f['text'] == 'Rust')
        self.assertEqual(entry['source'], 'editor')
        self.assertEqual(entry['dimension'], 'tech_stack')
        migrated = normalize_fact_store({'projects': [{'name': '旧原文', 'tech_stack': ['MCP']}]})
        self.assertTrue(supporting_ids('MCP', migrated['projects'][0]['facts']))

    def test_diff_version_restore_and_serializer(self):
        old, data = document(), document()
        data['projects'][0]['tech_stack'] = ['MCP']
        self.assertEqual(document_diff(old, data)[0]['after']['tech_stack'], ['MCP'])
        updated = prepare_version(old, {}, [], data, {})
        restored = restore_payload({**updated, 'collected_info': collected()})
        self.assertNotIn('tech_stack', restored['resume_data']['projects'][0])
        projected = graph._serialize_to_resume_data({**graph._init_state(), 'collected_info': collected(), 'resume_data': data})
        self.assertEqual(projected['projects'][0]['tech_stack'], ['MCP'])

    def test_pure_bullet_moves_with_coverage_and_citations(self):
        item = {'item_id': P, 'polished_bullets': ['技术栈：mcp、longchain', '实现导出'], 'description': ''}
        facts = [fact('f1', '技术栈：mcp、longchain'), fact('f2', '实现导出', kind='action')]
        meta = {'bullets': {P: [{'fact_ids': ['f1']}, {'fact_ids': ['f2']}]}}
        project_item_stack(item, facts, meta)
        self.assertEqual(item['tech_stack'], ['MCP', 'LangChain'])
        self.assertEqual(item['polished_bullets'], ['实现导出'])
        self.assertEqual(meta['bullets'][P][0]['fact_ids'], ['f2'])
        self.assertEqual(meta['tech_stack'][P][0]['fact_ids'], ['f1'])
        coverage = ensure_fact_coverage({'projects': [item]}, meta, {'projects': [{'item_id': P, 'facts': facts}]})
        self.assertEqual(next(row for row in coverage['facts'] if row['fact_id'] == 'f1')['status'], 'covered')

    def test_description_list_moves(self):
        item, meta = {'item_id': P, 'description': '技术栈：MCP、LangChain\n实现导出'}, {}
        project_item_stack(item, collected()['projects'][0]['facts'], meta, autofill=False)
        self.assertEqual(item['description'], '实现导出')
        self.assertEqual(item['tech_stack'], ['MCP', 'LangChain'])

    def test_ambiguous_removal_does_not_touch_multiple_items(self):
        data = {'projects': [{'name': name, 'tech_stack': ['MCP']} for name in ('甲', '乙')]}
        self.assertIn('error', parse_stack_edit(data, '把 MCP 去掉'))
        self.assertEqual(parse_stack_edit(data, '给甲项目加上技术栈 Python')['index'], 0)

    def test_unrecorded_edit_cannot_add(self):
        data = document()
        edit = parse_stack_edit(data, '给胡话简历项目加上技术栈 Rust')
        with self.assertRaisesRegex(ValueError, 'record_facts_first'):
            apply_stack_edit(data, {}, collected(), edit)

    def test_deleted_fact_cannot_restore_stack(self):
        data, store, meta = document(), collected(), {}
        project_item_stack(data['projects'][0], store['projects'][0]['facts'], meta)
        version = prepare_version(data, meta, [], document(), {})
        store['projects'][0]['facts'][0]['status'] = 'deleted'
        restored = restore_payload({**version, 'collected_info': store})
        self.assertFalse(restored['resume_data']['projects'][0].get('tech_stack'))


class TechGuardWriterTests(unittest.IsolatedAsyncioTestCase):
    def setUp(self):
        self.enterContext(patch('app.agents.tailor_pipeline.measure_rendered_pages', side_effect=measured))
        self.enterContext(patch('app.services.resume_capacity.measure_rendered_pages', side_effect=measured))

    async def test_guard_rejects_unknown_uncited_wrong_item_or_inactive(self):
        facts = [{**f, 'item_id': P} for f in collected()['projects'][0]['facts']]
        rows = [{'id': '/projects/0/tech_stack/0', 'text': 'InventedCloud', 'fact_ids': ['f1']},
                {'id': '/projects/0/tech_stack/1', 'text': 'MCP', 'fact_ids': []},
                {'id': 'wrong', 'text': 'MCP', 'claim_kind': 'tech_stack', 'item_id': W, 'fact_ids': ['f1']},
                {'id': 'stale', 'text': 'MCP', 'claim_kind': 'tech_stack', 'fact_ids': ['gone']}]
        with patch('app.services.fact_guard.verify_semantics', side_effect=AssertionError('no model')):
            checked = await guard_bullets(rows, facts)
        self.assertTrue(all(r['verdict'] == 'fail' for r in checked))
        self.assertIn('unsupported_tech:InventedCloud', checked[0]['lexical_errors'])

    async def test_guard_accepts_only_original_aliases(self):
        facts = collected()['projects'][0]['facts']
        rows = [{'id': f'/projects/0/tech_stack/{i}', 'text': name, 'fact_ids': ['f1']}
                for i, name in enumerate(['MCP', 'Tool Calling', 'LangChain'])]
        checked = await guard_bullets(rows, facts)
        self.assertTrue(all(r['verdict'] == 'pass' for r in checked))
        self.assertEqual((await guard_bullets([{'id': '/projects/0/tech_stack/0', 'text': 'MCP', 'fact_ids': ['x']}],
            [fact('x', 'MCP', source_quote='实现简历导出')]))[0]['verdict'], 'fail')

    async def test_legacy_writer_proposal_and_fallback(self):
        for propose in (True, False):
            result = await run_resume_writer(collected(), {'role_family': 'backend'}, WriterStub(propose))
            self.assertEqual(result['resume_data']['projects'][0]['tech_stack'], ['MCP', 'Tool Calling', 'LangChain', '长期记忆'])
            self.assertEqual(result['resume_data']['work_experience'][0]['tech_stack'], ['Python', 'FastAPI'])
            self.assertNotIn('InventedCloud', str(result['resume_data']))
            self.assertFalse(any(line.startswith('技术栈：') for line in result['resume_data']['work_experience'][0]['bullets']))

    async def test_labeled_writer_ends_with_separate_stack(self):
        result = await run_resume_writer(collected(), {'role_family': 'ai_pm'}, WriterStub(),
            previous_plan={'bullet_style': 'labeled'})
        self.assertIn('MCP', result['resume_data']['projects'][0]['tech_stack'])
        self.assertFalse(any(pure_stack(line) for line in result['resume_data']['projects'][0]['polished_bullets']))

    async def test_loop_write_uses_same_source_field(self):
        state = {**graph._init_state(), 'collected_info': collected(),
                 'messages': [HumanMessage(content='标一下技术栈并生成', id='draft')],
                 'confirmed_plan': {'bullet_style': 'plain', 'experience_plan': [], 'section_order': ['projects']}}
        result = await tools.write_draft(state, {}, tools.ToolContext(llm=WriterStub()))
        self.assertEqual(result.result['status'], 'ok', result.result)
        self.assertEqual(result.update['resume_data']['projects'][0]['tech_stack'], ['MCP', 'Tool Calling', 'LangChain', '长期记忆'])
        self.assertIn(P, result.update['resume_meta']['tech_stack'])

    async def test_loop_add_records_original_then_remove_and_undo(self):
        state = {**graph._init_state(), 'collected_info': collected(), 'resume_data': document(),
                 'messages': [HumanMessage(content='给胡话简历项目加上技术栈 Rust、longchain', id='add')]}
        async def intake(state, store):
            return {'collected_info': store}
        with patch.object(graph, '_persist_intake', side_effect=intake):
            result = await tools.edit_resume(state, {'instruction': '给项目添加 Kubernetes'}, tools.ToolContext())
        self.assertEqual(result.result['status'], 'ok', result.result)
        self.assertEqual(set(result.update['resume_data']['projects'][0]['tech_stack']), {'Rust', 'LangChain'})
        added = [f for f in result.update['collected_info']['projects'][0]['facts'] if f['id'] not in {'f1', 'f2'}]
        self.assertEqual([f['source_quote'] for f in added], ['给胡话简历项目加上技术栈 Rust、longchain'])
        current = {**state, **result.update, 'messages': [HumanMessage(content='把 LangChain 去掉', id='remove')]}
        removed = await tools.edit_resume(current, {'instruction': '把 LangChain 去掉'}, tools.ToolContext())
        self.assertEqual(removed.update['resume_data']['projects'][0]['tech_stack'], ['Rust'])
        self.assertEqual(removed.update['collected_info'], current['collected_info'])
        restored = restore_payload({**current, **removed.update})
        self.assertIn('LangChain', restored['resume_data']['projects'][0]['tech_stack'])
        rewritten = await run_resume_writer(removed.update['collected_info'], {'role_family': 'backend'}, WriterStub(),
            current_data=removed.update['resume_data'], resume_meta=removed.update['resume_meta'])
        self.assertNotIn('LangChain', rewritten['resume_data']['projects'][0]['tech_stack'])

    async def test_legacy_literal_stack_edit_path(self):
        state = {**graph._init_state(), 'collected_info': collected(), 'resume_data': document(),
                 'messages': [HumanMessage(content='给胡话简历项目加上技术栈 MCP', id='edit')]}
        with patch.object(graph, '_persist_intake', new=AsyncMock(side_effect=lambda state, store: {'collected_info': store})):
            result = await graph._evaluating_turn_node(state, WriterStub())
        self.assertEqual(result['resume_data']['projects'][0]['tech_stack'], ['MCP'])

    async def test_mcp_check_facts_and_write_gate(self):
        service = MCPService.__new__(MCPService)
        service._guard_facts = lambda principal: collected()['projects'][0]['facts']
        data = {'projects': [{'tech_stack': ['MCP', 'InventedCloud']}]}
        claims = resume_claims(data, {'/projects/0/tech_stack': ['f1']})
        result = await service.check_facts(None, claims)
        self.assertEqual(result['issue_count'], 1)
        with self.assertRaisesRegex(ValueError, 'unsupported_tech_stack'):
            await service._check_resume(None, data, {'/projects/0/tech_stack': ['f1']})
        self.assertIn('只填用户说过的技术', str(writing_guide()['expression_rules']))


class TechRenderTests(unittest.TestCase):
    def data(self):
        data = document()
        for section in ('projects', 'work_experience'):
            data[section][0]['tech_stack'] = ['MCP', 'Tool Calling', 'LangChain', '长期记忆']
        return data

    def test_all_existing_templates_render_both_paths_before_bullets(self):
        templates = list(TEMPLATE_DIR.glob('t*.html.j2'))
        self.assertGreaterEqual(len(templates), 9)
        for template in templates:
            with self.subTest(template=template.name):
                html = render_resume(template.name.removesuffix('.html.j2'), self.data())
                for section in ('projects', 'work_experience'):
                    self.assertIn(f'data-field-path="/{section}/0/tech_stack"', html)
                self.assertEqual(html.count('<strong>技术栈：</strong>MCP、Tool Calling、LangChain、长期记忆'), 2)
                self.assertLess(html.index('/work_experience/0/tech_stack'), html.index('整理接口文档并交付核对表'))

    def test_empty_and_hidden_never_render_tech_line(self):
        for template in TEMPLATE_DIR.glob('t*.html.j2'):
            for data in (document(), self.data()):
                if data['projects'][0].get('tech_stack'):
                    for section in ('projects', 'work_experience'):
                        data[section][0]['hidden'] = True
                self.assertNotIn('data-field-path="/projects/0/tech_stack"', render_resume(template.name.removesuffix('.html.j2'), data))

    def test_docx_line_order_and_empty(self):
        paras = [p.text for p in Document(io.BytesIO(resume_data_to_docx(self.data()))).paragraphs]
        tech = '技术栈：MCP、Tool Calling、LangChain、长期记忆'
        self.assertEqual(paras.count(tech), 2)
        self.assertEqual(paras[paras.index('实现简历导出并整理核对记录') - 1], tech)
        self.assertFalse(any('技术栈：' in p.text for p in Document(io.BytesIO(resume_data_to_docx(document()))).paragraphs))

    def test_real_pdf_measure_includes_line_height(self):
        from app.services.resume_pdf import measure_resume_pdf_layout
        from app.services.resume_capacity import resume_data_hash
        plain, stack = document(), self.data()
        # Real layout diagnostics expose the rendered field and its occupied height.
        a = measure_resume_pdf_layout('t001-jianyue', plain, 1, include_diagnostics=True)
        b = measure_resume_pdf_layout('t001-jianyue', stack, 1, include_diagnostics=True)
        self.assertNotEqual(resume_data_hash(plain), resume_data_hash(stack))
        self.assertGreater(b['last_page_content_height_px'], a['last_page_content_height_px'])
        before = {r['path']: r['lines'] for r in a['diagnostics']['experiences']}
        after = {r['path']: r['lines'] for r in b['diagnostics']['experiences']}
        for path in ('/projects/0', '/work_experience/0'):
            self.assertEqual(after[path], before[path] + 1)


@unittest.skipUnless(shutil.which('node'), 'node is not installed')
class TechDOMTests(unittest.TestCase):
    def test_editor_and_change_card(self):
        result = subprocess.run(['node', str(ROOT / 'tests/batch18b.dom.js')], capture_output=True, text=True, timeout=20)
        self.assertEqual(result.returncode, 0, result.stdout + result.stderr)


import test_mcp_protocol as mcp_fixture  # noqa: E402
from app.models.models import ResumeSession  # noqa: E402


class TechMCPPersistenceTests(unittest.IsolatedAsyncioTestCase):
    asyncSetUp = mcp_fixture.MCPProtocolTests.asyncSetUp
    cleanup = mcp_fixture.MCPProtocolTests.cleanup
    client = mcp_fixture.MCPProtocolTests.client
    call = mcp_fixture.MCPProtocolTests.call
    add = mcp_fixture.MCPProtocolTests.add

    async def test_mcp_schema_check_save_preview_export_and_web_save(self):
        self.enterContext(patch('app.api.resume.get_graph', return_value=graph.build_graph(None, self.saver)))
        self.enterContext(patch('app.api.resume._llm_for_graph', return_value=None))
        source = self.add(category='project', subject='胡话简历', field='tech_stack', value='mcp、longchain')
        work_source = self.add(category='work', subject='合成公司', field='tech_stack', value='MCP')
        data = {'projects': [{'name': '胡话简历', 'tech_stack': ['MCP', 'LangChain']}],
                'work_experience': [{'company': '合成公司', 'tech_stack': ['MCP']}]}
        citations = {'/projects/0/tech_stack': [source['id']], '/work_experience/0/tech_stack': [work_source['id']]}
        async with self.client() as client:
            checked = await self.call(client, 'huhua_resume_check_facts', {'bullets': resume_claims(data, citations)})
            self.assertTrue(all(row['verdict'] == 'pass' for row in checked['results'] if '/tech_stack/' in row['id']))
            saved = await self.call(client, 'huhua_resume_save', {'resume': data, 'citations': citations})
            sid = saved['conversation_id']
            got = await self.call(client, 'huhua_resume_get', {'conversation_id': sid})
            self.assertEqual(got['untrusted_text']['resume']['projects'][0]['tech_stack'], ['MCP', 'LangChain'])
            preview = await client.call_tool('huhua_resume_preview', {'conversation_id': sid, 'template_id': 't001-jianyue'})
            self.assertFalse(preview.is_error, preview.content)
            exported = await self.call(client, 'huhua_resume_export', {'conversation_id': sid, 'template_id': 't001-jianyue', 'format': 'docx'})
            content = (await self.web.get(exported['download_url'])).content
            self.assertIn('技术栈：MCP、LangChain', [p.text for p in Document(io.BytesIO(content)).paragraphs])
        current = (await self.web.get('/api/resume/current', params={'session_id': sid})).json()
        edited = deepcopy(current['data'])
        edited['projects'][0]['tech_stack'] = ['MCP', 'Rust']
        response = await self.web.put('/api/resume/current', json={'session_id': sid, 'expected_revision': current['revision'], 'resume_data': edited})
        self.assertEqual(response.status_code, 200, response.text)
        persisted = (await self.web.get('/api/resume/current', params={'session_id': sid})).json()
        self.assertEqual(persisted['data']['projects'][0]['tech_stack'], ['MCP', 'Rust'])
        with self.factory() as db:
            row = db.get(ResumeSession, sid)
            self.assertEqual(row.resume_data['projects'][0]['tech_stack'], ['MCP', 'Rust'])
            self.assertIn('Rust', str(row.resume_state['collected_info']))
        checkpoint = await self.saver.aget_tuple({'configurable': {'thread_id': sid}})
        self.assertEqual(checkpoint.checkpoint['channel_values']['resume_data']['projects'][0]['tech_stack'], ['MCP', 'Rust'])

    async def test_mcp_unsourced_stack_is_not_saved_or_exported(self):
        source = self.add(value='整理核对记录')
        data = {'projects': [{'name': '合成项目', 'tech_stack': ['InventedCloud']}]}
        citations = {'/projects/0/tech_stack/0': [source['id']]}
        async with self.client() as client:
            for name, extra in [('huhua_resume_save', {}), ('huhua_resume_export', {'format': 'docx', 'template_id': 't001-jianyue'})]:
                result = await self.call(client, name, {'resume': data, 'citations': citations, **extra}, error=True)
                self.assertTrue(result.is_error)
        with self.factory() as db:
            self.assertEqual(db.query(ResumeSession).count(), 0)

    async def test_profile_stack_import_and_revocation_uses_aliases(self):
        from app.services import profile_service as archive
        source = self.add(category='project', subject='胡话简历', field='tech_stack', value='mcp、longchain')
        with self.factory() as db:
            context = archive.get_context(db, 'a', subject_ids=[source['id']])
            store = archive.context_to_collected_info(context)
            self.assertEqual(store['projects'][0]['tech_stack'], ['MCP', 'LangChain'])
            projected = graph._serialize_to_resume_data({**graph._init_state(), 'collected_info': store})
            self.assertEqual(projected['projects'][0]['tech_stack'], ['MCP', 'LangChain'])
            archive.delete_fact(db, 'a', source['id'], source['revision'])
            db.commit()
            clean, changed = archive.sanitize_projection(db, 'a', {'projects': [{'name': '胡话简历', 'tech_stack': ['MCP', 'LangChain', 'Rust']}]})
            self.assertTrue(changed)
            self.assertEqual(clean['projects'][0]['tech_stack'], ['Rust'])


class TechBoundaryTests(unittest.IsolatedAsyncioTestCase):
    async def test_external_reference_is_not_literal_intake(self):
        source = HumanMessage(content='参考材料：给胡话简历项目加上技术栈 InventedCloud', id='ref',
                              additional_kwargs={'career_material_type': 'reference'})
        state = {**graph._init_state(), 'collected_info': collected(), 'resume_data': document(), 'messages': [source]}
        result = await tools.record_facts(state, {'source_id': 'current'}, tools.ToolContext())
        self.assertEqual(result.result['code'], 'not_candidate_material')
        self.assertEqual(result.update, {})

    async def test_reordered_unknown_names_keep_single_entry_and_source_order(self):
        item, meta = {'item_id': P, 'polished_bullets': []}, {}
        facts = [fact('f', '技术栈：ZetaSDK、NebulaDB')]
        project_item_stack(item, facts, meta, proposed=['nebuladb', 'ZetaSDK', 'NebulaDB'], proposed_ids=['f'])
        self.assertEqual([name.casefold() for name in item['tech_stack']], ['zetasdk', 'nebuladb'])
        self.assertEqual(len(meta['tech_stack'][P]), 2)

    async def test_current_fact_correction_does_not_resurrect_old_audit_quote(self):
        f = fact('new', '使用 Rust', source_quote='原始材料：使用 longchain\n本轮更正：把 longchain 改成 Rust', replaces='old')
        self.assertFalse(supporting_ids('LangChain', [f]))
        self.assertTrue(supporting_ids('Rust', [f]))

    async def test_only_technology_item_stays_bullet_free_after_loop_fallback(self):
        store = {'projects': [{'item_id': P, 'name': '胡话简历', 'facts': [fact('f', '技术栈：mcp、longchain')]}]}
        state = {**graph._init_state(), 'collected_info': store,
                 'messages': [HumanMessage(content='生成简历', id='draft')],
                 'confirmed_plan': {'bullet_style': 'plain', 'experience_plan': [], 'section_order': ['projects']}}
        with patch('app.agents.tailor_pipeline.measure_rendered_pages', side_effect=measured), patch('app.services.resume_capacity.measure_rendered_pages', side_effect=measured):
            result = await tools.write_draft(state, {}, tools.ToolContext(llm=WriterStub()))
        self.assertEqual(result.result['status'], 'ok', result.result)
        item = result.update['resume_data']['projects'][0]
        self.assertEqual(item['tech_stack'], ['MCP', 'LangChain'])
        self.assertEqual(item.get('polished_bullets', []), [])

    async def test_work_literal_add_remove(self):
        data = document()
        add = parse_stack_edit(data, '给合成公司工作经历加上技术栈 Python、FastAPI')
        data, meta = apply_stack_edit(data, {}, collected(), add)
        self.assertEqual(data['work_experience'][0]['tech_stack'], ['Python', 'FastAPI'])
        remove = parse_stack_edit(data, '合成公司把 FastAPI 去掉')
        data, meta = apply_stack_edit(data, meta, collected(), remove)
        self.assertEqual(data['work_experience'][0]['tech_stack'], ['Python'])
        version = graph._serialize_to_resume_data({'collected_info': collected(), 'resume_data': data})
        self.assertEqual(version['work_experience'][0]['tech_stack'], ['Python'])


class TechStackProseBoundaryTests(unittest.TestCase):
    def test_space_separated_list_splits_known_names(self):
        for quote in ('我们这个项目使用了 mcp 记忆 toolcalling longchain',
                      '这个项目用了 mcp 记忆 tool calling longchain'):
            self.assertEqual(extract_stack({'source_quote': quote}), ['MCP', '记忆', 'Tool Calling', 'LangChain'])

    def test_multiword_name_without_known_token_stays_whole(self):
        self.assertEqual(extract_stack({'source_quote': '技术栈：Spring Boot、Redis'}), ['Spring Boot', 'Redis'])

    def test_chinese_capability_word_in_prose_is_not_stack(self):
        self.assertEqual(extract_stack({'source_quote': '做了一个背单词小程序，帮用户记忆单词，用 Python 写后端'}), ['Python'])
