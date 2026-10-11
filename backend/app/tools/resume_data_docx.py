"""
从 ResumeData dict 生成 Word 文档（不走 markdown 中间层）。

与 docx_exporter.py 的区别：
  - docx_exporter: 输入聊天产出的 markdown 文本（resume_draft）
  - 本模块:        输入结构化 ResumeData（前端表单 / 聊天 state.resume_data）
"""

import base64
import io
import logging
from typing import Any

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Inches, Mm, Pt

from app.services.resume_sections import get_resume_sections, visible_items
from app.services.resume_field_hygiene import award_label, contribution_title
from app.services.resume_layout import layout_css_values, resume_density

logger = logging.getLogger(__name__)


def _decode_photo_data_url(data_url: str) -> bytes | None:
    """解析 data:image/...;base64,xxx 成 bytes。格式错返回 None。"""
    if not data_url or not data_url.startswith("data:image/"):
        return None
    try:
        _, b64 = data_url.split(",", 1)
        return base64.b64decode(b64)
    except Exception as e:
        logger.warning(f"[docx/photo] 解码失败: {e}")
        return None


def _add_section_heading(doc, text: str, data: dict) -> None:
    if not text:
        return
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = True
    values = layout_css_values('t001-jianyue', data)
    run.font.size = Pt(float(values['section-font-size'][:-2]) if resume_density(data) else 13)
    if resume_density(data):
        p.paragraph_format.space_before = Pt(float(values['section-top'][:-2]) * .75)
        p.paragraph_format.space_after = Pt(float(values['section-bottom'][:-2]) * .75)


def _add_kv_line(doc, label: str, value: str) -> None:
    if not value:
        return
    p = doc.add_paragraph()
    run = p.add_run(f"{label}：{value}")
    run.font.size = Pt(10.5)


def resume_data_to_docx(data: dict[str, Any]) -> bytes:
    """把 ResumeData dict 转为 docx bytes。空字段不输出对应行。"""
    doc = Document()
    for section in doc.sections:
        section.page_width = Mm(210)
        section.page_height = Mm(297)
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
    style = doc.styles["Normal"]
    style.font.name = "微软雅黑"
    # python-docx's font.name only sets Latin mappings; map CJK explicitly too.
    style.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "微软雅黑")
    style.font.size = Pt(10.5)

    basic = data.get("basic_info") or {}

    # 头像（有则嵌入；doc.add_picture 放 inline 段落，右对齐看视觉更像简历照片位）
    photo_bytes = _decode_photo_data_url(basic.get("photo") or "")
    if photo_bytes:
        try:
            photo_p = doc.add_paragraph()
            run = photo_p.add_run()
            run.add_picture(io.BytesIO(photo_bytes), width=Inches(1.2))
            # 右对齐照片段
            from docx.enum.text import WD_ALIGN_PARAGRAPH

            photo_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        except Exception as e:
            logger.warning(f"[docx/photo] add_picture 失败: {e}")

    # 标题：姓名（大号）
    if basic.get("name"):
        title = doc.add_paragraph()
        run = title.add_run(basic["name"])
        run.bold = True
        run.font.size = Pt(18)

    # 基本信息一行拼起来
    meta_parts = []
    for k in ("phone", "email", "location", "objective"):
        if basic.get(k):
            meta_parts.append(basic[k])
    if meta_parts:
        p = doc.add_paragraph()
        p.add_run(" · ".join(meta_parts))

    # 额外 basic_info：身高/年龄/政治面貌/民族（仅显示非空）
    opt_parts = []
    opt_labels = {
        "height": "身高",
        "age": "年龄",
        "gender": "性别",
        "political": "政治面貌",
        "ethnicity": "民族",
        "hometown": "籍贯",
        "birthday": "出生年月",
    }
    for k, label in opt_labels.items():
        if basic.get(k):
            opt_parts.append(f"{label} {basic[k]}")
    if opt_parts:
        doc.add_paragraph(" | ".join(opt_parts))
    if basic.get("availability"):
        doc.add_paragraph("可实习时间：" + basic["availability"])
    if basic.get('links'):
        p = doc.add_paragraph()
        for index, link in enumerate(basic['links']):
            if index:
                p.add_run(' · ')
            _add_link(p, link)
    if data.get("summary"):
        doc.add_paragraph("个人定位").runs[0].bold = True
        doc.add_paragraph(data["summary"])

    # 预览与 Word 共享栏目顺序、标题与可见性；正文仍取原有字段。
    for resume_section in get_resume_sections(data):
        kind = resume_section["kind"]
        if kind == "custom":
            has_content = bool(resume_section["title"] or resume_section["content"])
        elif kind == "skills":
            has_content = any((data.get("skill_groups") or {}).values()) or bool(data.get("skills"))
        elif kind in {"education", "work_experience", "projects", "open_source", "awards"}:
            has_content = bool(visible_items(data, kind))
        else:
            has_content = bool(data.get(kind))
        if not has_content:
            continue
        _add_section_heading(doc, resume_section["title"], data)
        _add_section_body(doc, data, resume_section)

    configure_docx_layout(doc, data)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _add_section_body(doc, data: dict[str, Any], section: dict[str, Any]) -> None:
    kind = section["kind"]
    if kind == "education":
        for _, edu in visible_items(data, "education"):
            line = " / ".join(
                x
                for x in [
                    edu.get("school"),
                    edu.get("major"),
                    edu.get("degree"),
                    f"{edu.get('start_date', '')}—{edu.get('end_date', '')}".strip("—"),
                ]
                if x and x != "—"
            )
            p = doc.add_paragraph()
            p.add_run(line).bold = True
            if edu.get("highlights"):
                doc.add_paragraph(edu["highlights"])

    elif kind == "work_experience":
        for _, job in visible_items(data, "work_experience"):
            line = " / ".join(
                x
                for x in [
                    job.get("company"),
                    job.get("title"),
                    f"{job.get('start_date', '')}—{job.get('end_date', '')}".strip("—"),
                ]
                if x and x != "—"
            )
            p = doc.add_paragraph()
            p.add_run(line).bold = True
            for b in job.get("bullets") or []:
                doc.add_paragraph(b, style="List Bullet")
    elif kind == "projects":
        for _, proj in visible_items(data, "projects"):
            line = " / ".join(
                x
                for x in [
                    proj.get("name"),
                    proj.get("role"),
                    f"{proj.get('start_date', '')}—{proj.get('end_date', '')}".strip(
                        "—"
                    ),
                ]
                if x and x != "—"
            )
            p = doc.add_paragraph()
            p.add_run(line).bold = True
            bullets = proj.get("polished_bullets") or []
            if bullets:
                for b in bullets:
                    doc.add_paragraph(b, style="List Bullet")
            elif proj.get("description"):
                doc.add_paragraph(proj["description"])

    elif kind in {"open_source", "awards"}:
        items = visible_items(data, kind)
        if kind == "awards":
            p = doc.add_paragraph("；".join(award_label(item) for _, item in items))
            p.paragraph_format.space_after = Pt(3)
        for _, item in items:
            if kind == "open_source":
                date = f"{item.get('start_date', '')}—{item.get('end_date', '')}".strip("—")
                title = doc.add_paragraph()
                title.add_run(contribution_title(item) + (" / " + date if date else "")).bold = True
                for link in item.get('links') or []:
                    title.add_run(' · ')
                    _add_link(title, link, label='链接')
            for _, bullet in item["_display_bullets"]:
                doc.add_paragraph((item["name"] + "：" if kind == "awards" and len(items) > 1 else "") + bullet, style="List Bullet")
            for link in (item.get("links") or []) if kind == 'awards' else []:
                doc.add_paragraph(f"{link.get('text', '')}：{link.get('url', '')}")

    elif kind == "skills":
        # 分组技能优先，fallback 到 flat skills。
        skill_groups = data.get("skill_groups") or {}
        if skill_groups and any(skill_groups.values()):
            from app.services.skill_groups import display_skill_rows
            for row in display_skill_rows(skill_groups):
                p = doc.add_paragraph()
                for index, (group_name, items) in enumerate(row):
                    if index:
                        p.add_run('；')
                    p.add_run(f"{group_name}：").bold = True
                    p.add_run(" / ".join(items))
        else:
            for skill in data.get("skills") or []:
                doc.add_paragraph(skill, style="List Bullet")

    elif kind == "certificates":
        for c in data["certificates"]:
            doc.add_paragraph(c, style="List Bullet")
    elif kind == "self_evaluation":
        doc.add_paragraph(data["self_evaluation"])
    elif kind == "custom":
        # add_paragraph treats HTML as literal text; each authored line is retained.
        if section["content"]:
            for line in section["content"].replace("\r\n", "\n").replace("\r", "\n").split("\n"):
                doc.add_paragraph(line)


def _add_link(paragraph, link: dict, *, label: str = '') -> None:
    from docx.oxml import OxmlElement
    from docx.opc.constants import RELATIONSHIP_TYPE
    from app.services.renderer import safe_link_url

    url = str(link.get('url') or '')
    safe = safe_link_url(url)
    text = label or f"{link.get('text', '')}：{url if len(url) <= 60 else '链接'}"
    if not safe:
        paragraph.add_run(f"{link.get('text', '')}：{url}")
        return
    node = OxmlElement('w:hyperlink')
    node.set(qn('r:id'), paragraph.part.relate_to(safe, RELATIONSHIP_TYPE.HYPERLINK, is_external=True))
    run = OxmlElement('w:r')
    value = OxmlElement('w:t')
    value.text = text
    run.append(value)
    node.append(run)
    paragraph._p.append(node)


def configure_docx_layout(doc, data: dict) -> None:
    """Apply the same density to generic Word and legacy filled templates."""
    density = resume_density(data)
    doc.core_properties.keywords = f'resume-density={density}'
    if not density:
        return
    values = layout_css_values('t001-jianyue', data)
    font = float(values['body-font-size'][:-2])
    # Word's built-in default is already 1.15, unlike the HTML templates.
    # Tightening must not expand Word to the HTML density-one line height.
    line = (1.15, 1.1, 1.05, 1.0)[density]
    for name in ('Normal', 'List Bullet'):
        style = doc.styles[name]
        style.font.size = Pt(font if name == 'Normal' else float(values['content-font-size'][:-2]))
        style.paragraph_format.line_spacing = line
        style.paragraph_format.space_before = Pt(0)
        style.paragraph_format.space_after = Pt(float(values['li-gap' if name == 'List Bullet' else 'item-gap'][:-2]) * .75)
    factor = (1, 1, .92, .88)[density]
    for section in doc.sections:
        section.top_margin = int(section.top_margin * factor)
        section.bottom_margin = int(section.bottom_margin * factor)
        section.left_margin = int(section.left_margin * factor)
        section.right_margin = int(section.right_margin * factor)
    # Legacy templates have direct formatting in table cells and runs. Normalize
    # those overrides as well, retaining large identity/heading text.
    def paragraphs(parent):
        yield from parent.paragraphs
        for table in parent.tables:
            for row in table.rows:
                for cell in row.cells:
                    yield from paragraphs(cell)
    headings = {section['title'] for section in get_resume_sections(data)}
    for paragraph in paragraphs(doc):
        previous_line = paragraph.paragraph_format.line_spacing
        paragraph.paragraph_format.line_spacing = min(previous_line, line) if isinstance(previous_line, float) else line
        if paragraph.text in headings:
            continue
        for name in ('space_before', 'space_after'):
            existing = getattr(paragraph.paragraph_format, name)
            if existing is not None:
                setattr(paragraph.paragraph_format, name, int(existing * (1, .65, .5, .4)[density]))
        for run in paragraph.runs:
            if run.font.size is not None and run.font.size.pt <= 11:
                run.font.size = Pt(font)
